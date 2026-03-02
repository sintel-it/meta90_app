from datetime import datetime, timedelta
import calendar as pycalendar
import csv
import json
import os
import secrets
import sqlite3
import re
import base64
import hashlib
import sys
import subprocess
import zipfile
from io import StringIO
from io import BytesIO
from urllib import error as urlerror
from urllib import parse as urlparse
from urllib import request as urlrequest

from flask import Flask, Response, abort, flash, has_request_context, jsonify, make_response, redirect, render_template, request, session, url_for
from itsdangerous import BadSignature, URLSafeSerializer
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

from migrations import apply_migrations
from modules import calendario_logic, calendario_queries, mensajes_logic, metas_logic
from modules import notificaciones_compose
from routes.auth import auth_bp, configure_auth_routes
from routes.calendario import calendario_bp, configure_calendario_routes
from routes.mensajes import configure_mensajes_routes, mensajes_bp
from routes.metas import configure_metas_routes, metas_bp
from routes.notificaciones import configure_notificaciones_routes, notificaciones_bp
from routes.perfil import configure_perfil_routes, perfil_bp
from services import mail_api, push_api, sms_api, whatsapp_api


def cargar_env_local(ruta=".env"):
    if not os.path.exists(ruta):
        return

    with open(ruta, "r", encoding="utf-8") as archivo:
        for linea in archivo:
            linea = linea.strip()
            if not linea or linea.startswith("#") or "=" not in linea:
                continue

            clave, valor = linea.split("=", 1)
            clave = clave.strip()
            valor = valor.strip().strip('"').strip("'")

            # Hardening: variables ya definidas por el sistema no se sobreescriben.
            if clave and os.getenv(clave) in (None, ""):
                os.environ[clave] = valor


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
cargar_env_local(os.path.join(BASE_DIR, ".env"))
MAINTENANCE_NOTIFICATIONS_FLAG = os.path.join(BASE_DIR, "logs", "maintenance_notificaciones.flag")

app = Flask(__name__)
APP_ENV = os.getenv("APP_ENV", "").strip().lower()
IS_PROD = APP_ENV in ("prod", "production") or bool(os.getenv("RENDER"))
_secret_from_env = os.getenv("FLASK_SECRET_KEY", "").strip() or os.getenv("SECRET_KEY", "").strip()
app.secret_key = _secret_from_env or secrets.token_hex(32)
if IS_PROD and not _secret_from_env:
    print("WARN: SECRET_KEY/FLASK_SECRET_KEY no definido en produccion; se usa clave temporal.", file=sys.stderr)
SESSION_DAYS = max(1, int(os.getenv("SESSION_DAYS", "30")))
REMEMBER_ME_DAYS = max(1, int(os.getenv("REMEMBER_ME_DAYS", "90")))
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "1" if IS_PROD else "0").strip() in ("1", "true", "True")
COOKIE_SAMESITE = os.getenv("COOKIE_SAMESITE", "Lax").strip() or "Lax"
app.permanent_session_lifetime = timedelta(days=SESSION_DAYS)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SECURE"] = COOKIE_SECURE
app.config["SESSION_COOKIE_SAMESITE"] = COOKIE_SAMESITE
# Render/Reverse proxy support for correct scheme/host on OAuth callbacks.
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_port=1)
DB_PATH = os.getenv("DB_PATH", "").strip() or os.path.join(BASE_DIR, "metas.db")
BACKUP_DIR = os.getenv("BACKUP_DIR", "").strip() or os.path.join(BASE_DIR, "backups")

SMTP_HOST = os.getenv("SMTP_HOST", "").strip()
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "").strip()
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "")
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER).strip()
SMTP_USE_TLS = os.getenv("SMTP_USE_TLS", "1").strip() not in ("0", "false", "False")

RESEND_API_KEY = os.getenv("RESEND_API_KEY", "").strip()
RESEND_FROM = os.getenv("RESEND_FROM", "").strip()
MAIL_PROVIDER = os.getenv("MAIL_PROVIDER", "auto").strip().lower()
FACEBOOK_APP_ID = os.getenv("FACEBOOK_APP_ID", "").strip()
FACEBOOK_APP_SECRET = os.getenv("FACEBOOK_APP_SECRET", "").strip()
FACEBOOK_SCOPE = os.getenv("FACEBOOK_SCOPE", "email,public_profile").strip()
FACEBOOK_GRAPH_VERSION = os.getenv("FACEBOOK_GRAPH_VERSION", "v21.0").strip()
GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "").strip()
GOOGLE_CLIENT_SECRET = os.getenv("GOOGLE_CLIENT_SECRET", "").strip()
GOOGLE_SCOPE = os.getenv("GOOGLE_SCOPE", "openid email profile").strip()
GOOGLE_REDIRECT_URI = os.getenv("GOOGLE_REDIRECT_URI", "").strip()
MICROSOFT_CLIENT_ID = os.getenv("MICROSOFT_CLIENT_ID", "").strip()
MICROSOFT_CLIENT_SECRET = os.getenv("MICROSOFT_CLIENT_SECRET", "").strip()
MICROSOFT_TENANT = os.getenv("MICROSOFT_TENANT", "common").strip()
MICROSOFT_SCOPE = os.getenv("MICROSOFT_SCOPE", "openid profile email User.Read").strip()
ADMIN_2FA_REQUIRED = os.getenv("ADMIN_2FA_REQUIRED", "1").strip() in ("1", "true", "True")

WHATSAPP_TOKEN = os.getenv("WHATSAPP_TOKEN", "").strip()
WHATSAPP_PHONE_NUMBER_ID = os.getenv("WHATSAPP_PHONE_NUMBER_ID", "").strip()
WHATSAPP_API_VERSION = os.getenv("WHATSAPP_API_VERSION", "v21.0").strip()
WHATSAPP_NOTI_MODE = os.getenv("WHATSAPP_NOTI_MODE", "template_then_text").strip().lower()
WHATSAPP_TEMPLATE_NAME = os.getenv("WHATSAPP_TEMPLATE_NAME", "hello_world").strip()
WHATSAPP_TEMPLATE_LANG = os.getenv("WHATSAPP_TEMPLATE_LANG", "en_US").strip()
TWILIO_ACCOUNT_SID = os.getenv("TWILIO_ACCOUNT_SID", "").strip()
TWILIO_AUTH_TOKEN = os.getenv("TWILIO_AUTH_TOKEN", "").strip()
TWILIO_FROM_NUMBER = os.getenv("TWILIO_FROM_NUMBER", "").strip()
SMS_DEFAULT_COUNTRY_CODE = os.getenv("SMS_DEFAULT_COUNTRY_CODE", "+1").strip() or "+1"
ENABLE_TEST_NOTIFICATIONS = os.getenv("ENABLE_TEST_NOTIFICATIONS", "0").strip() in ("1", "true", "True")
SMS_TEST_MIN_NOTIFICATIONS = max(0, int(os.getenv("SMS_TEST_MIN_NOTIFICATIONS", "0"))) if ENABLE_TEST_NOTIFICATIONS else 0
FORCE_NOTIFICATIONS_TEST = (
    ENABLE_TEST_NOTIFICATIONS
    and os.getenv("FORCE_NOTIFICATIONS_TEST", "0").strip() in ("1", "true", "True")
)
PUSH_VAPID_PUBLIC_KEY = os.getenv("PUSH_VAPID_PUBLIC_KEY", "").strip()
PUSH_VAPID_PRIVATE_KEY = os.getenv("PUSH_VAPID_PRIVATE_KEY", "").strip()
PUSH_VAPID_SUBJECT = os.getenv("PUSH_VAPID_SUBJECT", "mailto:admin@meta.local").strip()
ALERT_WEBHOOK_URL = os.getenv("ALERT_WEBHOOK_URL", "").strip()
BACKUP_OFFSITE_DIR = os.getenv("BACKUP_OFFSITE_DIR", "").strip()
RATE_LIMIT_CONFIG = {
    "login": {"max_attempts": 5, "window_seconds": 300, "block_seconds": 600},
    "recuperar": {"max_attempts": 3, "window_seconds": 900, "block_seconds": 900},
    "registro_wa": {"max_attempts": 3, "window_seconds": 600, "block_seconds": 900},
    "api_push": {"max_attempts": 30, "window_seconds": 300, "block_seconds": 300},
    "api_admin": {"max_attempts": 20, "window_seconds": 300, "block_seconds": 300},
    "api_busqueda": {"max_attempts": 120, "window_seconds": 300, "block_seconds": 120},
}
MESES_ES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)
DIAS_ES_CORTO = ("lun.", "mar.", "mie.", "jue.", "vie.", "sab.", "dom.")
MODULES_WITH_PERMS = ("metas", "calendario", "mensajes", "notificaciones", "perfil")


def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _calcular_diff(antes, despues):
    if not isinstance(antes, dict) or not isinstance(despues, dict):
        return {}
    keys = set(antes.keys()) | set(despues.keys())
    diff = {}
    for k in sorted(keys):
        a = antes.get(k)
        d = despues.get(k)
        if a != d:
            diff[k] = {"antes": a, "despues": d}
    return diff


def registrar_auditoria(actor_user_id, modulo, accion, entidad, entidad_id=None, detalle=None, antes=None, despues=None):
    if not actor_user_id:
        return
    payload = dict(detalle or {})
    if has_request_context():
        payload["ip"] = obtener_ip_cliente()
        payload["user_agent"] = (request.headers.get("User-Agent", "") or "")[:220]
    if antes is not None and despues is not None:
        payload["diff"] = _calcular_diff(antes, despues)

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO audit_log
        (actor_user_id, modulo, accion, entidad, entidad_id, detalle, creado_en)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            int(actor_user_id),
            str(modulo or "")[:50],
            str(accion or "")[:50],
            str(entidad or "")[:50],
            int(entidad_id) if entidad_id is not None else None,
            json.dumps(payload, ensure_ascii=True)[:2000],
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    )
    conn.commit()
    conn.close()


def crear_base_datos():
    conn = get_connection()
    cursor = conn.cursor()
    apply_migrations(cursor)

    # Usuario admin por defecto.
    cursor.execute("SELECT id FROM usuarios WHERE username = ?", ("admin",))
    admin = cursor.fetchone()
    if admin is None:
        cursor.execute(
            "INSERT INTO usuarios (username, password, email, rol) VALUES (?, ?, ?, ?)",
            ("admin", generate_password_hash("1234"), "admin@meta.local", "admin"),
        )
        admin_id = cursor.lastrowid
    else:
        admin_id = admin["id"]
        cursor.execute("UPDATE usuarios SET rol = 'admin' WHERE id = ?", (admin_id,))

    # Asignar metas antiguas sin user_id al admin.
    cursor.execute("UPDATE metas SET user_id = ? WHERE user_id IS NULL", (admin_id,))
    # Migra contrasenas antiguas en texto plano a hash seguro.
    migrar_passwords_legacy(cursor)

    conn.commit()
    conn.close()


def usuario_autenticado():
    return "usuario" in session


def obtener_usuario_id_actual():
    user_id = session.get("usuario_id")
    if user_id:
        return user_id

    username = session.get("usuario")
    if not username:
        return None

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM usuarios WHERE username = ?", (username,))
    fila = cursor.fetchone()
    conn.close()

    if fila is None:
        return None

    session["usuario_id"] = fila["id"]
    return fila["id"]


def obtener_rol_usuario_actual():
    rol = (session.get("usuario_rol") or "").strip().lower()
    if rol in ("admin", "editor", "lector"):
        return rol

    user_id = obtener_usuario_id_actual()
    if user_id is None:
        return None

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT rol FROM usuarios WHERE id = ?", (user_id,))
    fila = cursor.fetchone()
    conn.close()

    rol_db = (fila["rol"] if fila and fila["rol"] else "editor").strip().lower() if fila else "editor"
    if rol_db not in ("admin", "editor", "lector"):
        rol_db = "editor"
    session["usuario_rol"] = rol_db
    return rol_db


def usuario_tiene_rol(rol):
    return obtener_rol_usuario_actual() == (rol or "").strip().lower()


def usuario_es_admin():
    return usuario_tiene_rol("admin")


def _default_module_perm_by_role(rol, modulo):
    rol = (rol or "").strip().lower()
    modulo = (modulo or "").strip().lower()
    if rol == "admin":
        return {"view": True, "edit": True}
    if rol == "lector":
        return {"view": True, "edit": False}
    # editor por defecto
    return {"view": True, "edit": True}


def obtener_permiso_modulo_usuario(user_id, modulo):
    if not user_id:
        return _default_module_perm_by_role("editor", modulo)
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT coalesce(rol, 'editor') AS rol FROM usuarios WHERE id = ?", (user_id,))
    user_row = cursor.fetchone()
    rol = (user_row["rol"] if user_row and user_row["rol"] else "editor").strip().lower() if user_row else "editor"
    base = _default_module_perm_by_role(rol, modulo)
    cursor.execute(
        """
        SELECT can_view, can_edit
        FROM user_module_permissions
        WHERE user_id = ? AND modulo = ?
        """,
        (user_id, modulo),
    )
    row = cursor.fetchone()
    conn.close()
    if not row:
        return base
    return {"view": bool(row["can_view"]), "edit": bool(row["can_edit"])}


def _modulo_from_endpoint(endpoint):
    ep = (endpoint or "").strip()
    mapping = {
        "metas_bp.": "metas",
        "calendario_bp.": "calendario",
        "mensajes_bp.": "mensajes",
        "notificaciones_bp.": "notificaciones",
        "perfil_bp.": "perfil",
    }
    for prefix, modulo in mapping.items():
        if ep.startswith(prefix):
            return modulo
    return None


def token_csrf_actual():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


def parsear_fecha_limite(fecha_texto):
    if not fecha_texto:
        return None

    fecha_texto = str(fecha_texto).strip()
    formatos = (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
    )

    for formato in formatos:
        try:
            return datetime.strptime(fecha_texto, formato).date()
        except ValueError:
            continue

    return None


def parsear_fecha_hora(fecha_texto):
    if not fecha_texto:
        return None
    try:
        return datetime.strptime(str(fecha_texto).strip(), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def hash_contrasena(contrasena):
    return generate_password_hash(contrasena)


def contrasena_parece_hash(contrasena_guardada):
    valor = (contrasena_guardada or "").strip().lower()
    return valor.startswith("scrypt:") or valor.startswith("pbkdf2:") or valor.startswith("argon2:")


def verificar_contrasena(contrasena_guardada, contrasena_ingresada):
    if not contrasena_guardada:
        return False, False

    if contrasena_guardada == contrasena_ingresada:
        # Compatibilidad con registros legacy en texto plano.
        return True, True

    try:
        return check_password_hash(contrasena_guardada, contrasena_ingresada), False
    except ValueError:
        return False, False


def migrar_passwords_legacy(cursor):
    cursor.execute("SELECT id, password FROM usuarios")
    usuarios = cursor.fetchall()

    for user in usuarios:
        password_actual = (user["password"] or "").strip()
        if not password_actual or contrasena_parece_hash(password_actual):
            continue

        cursor.execute(
            "UPDATE usuarios SET password = ? WHERE id = ?",
            (hash_contrasena(password_actual), user["id"]),
        )


def obtener_ip_cliente():
    forwarded_for = request.headers.get("X-Forwarded-For", "").strip()
    if forwarded_for:
        return forwarded_for.split(",")[0].strip()
    return (request.remote_addr or "desconocida").strip()


def _segundos_actuales():
    return int(datetime.now().timestamp())


def marca_envio_actual():
    ahora = datetime.now()
    franja = "manana" if ahora.hour < 15 else "noche"
    return f"{ahora.strftime('%Y-%m-%d')}|{franja}"


def _prefs_notificacion_usuario(user_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT allow_email, allow_sms, allow_push, morning_hour, night_hour, quiet_days
        FROM user_notification_prefs
        WHERE user_id = ?
        """,
        (user_id,),
    )
    row = cursor.fetchone()
    if row is None:
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute(
            """
            INSERT INTO user_notification_prefs
            (user_id, allow_email, allow_sms, allow_push, morning_hour, night_hour, quiet_days, actualizado_en)
            VALUES (?, 1, 1, 1, 8, 20, '', ?)
            """,
            (user_id, now),
        )
        conn.commit()
        cursor.execute(
            """
            SELECT allow_email, allow_sms, allow_push, morning_hour, night_hour, quiet_days
            FROM user_notification_prefs
            WHERE user_id = ?
            """,
            (user_id,),
        )
        row = cursor.fetchone()
    conn.close()
    return {
        "allow_email": int(row["allow_email"]) == 1,
        "allow_sms": int(row["allow_sms"]) == 1,
        "allow_push": int(row["allow_push"]) == 1,
        "morning_hour": int(row["morning_hour"] or 8),
        "night_hour": int(row["night_hour"] or 20),
        "quiet_days": str(row["quiet_days"] or ""),
    }


def obtener_prefs_notificacion_usuario(user_id):
    data = _prefs_notificacion_usuario(user_id)
    data["quiet_days_list"] = sorted(list(_quiet_days_set(data.get("quiet_days", ""))))
    return data


def guardar_prefs_notificacion_usuario(user_id, prefs):
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO user_notification_prefs
        (user_id, allow_email, allow_sms, allow_push, morning_hour, night_hour, quiet_days, actualizado_en)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(user_id) DO UPDATE SET
            allow_email=excluded.allow_email,
            allow_sms=excluded.allow_sms,
            allow_push=excluded.allow_push,
            morning_hour=excluded.morning_hour,
            night_hour=excluded.night_hour,
            quiet_days=excluded.quiet_days,
            actualizado_en=excluded.actualizado_en
        """,
        (
            user_id,
            int(prefs.get("allow_email", 1)),
            int(prefs.get("allow_sms", 1)),
            int(prefs.get("allow_push", 1)),
            int(prefs.get("morning_hour", 8)),
            int(prefs.get("night_hour", 20)),
            str(prefs.get("quiet_days", ""))[:64],
            now,
        ),
    )
    conn.commit()
    conn.close()


def _quiet_days_set(text):
    out = set()
    for p in str(text or "").split(","):
        p = p.strip()
        if p.isdigit():
            v = int(p)
            if 0 <= v <= 6:
                out.add(v)
    return out


def _canal_notificacion_habilitado(user_id, canal, forzar_envio=False):
    prefs = _prefs_notificacion_usuario(user_id)
    if canal == "email" and not prefs["allow_email"]:
        return False, "Canal email desactivado en perfil."
    if canal == "sms" and not prefs["allow_sms"]:
        return False, "Canal SMS desactivado en perfil."
    if canal == "push" and not prefs["allow_push"]:
        return False, "Canal push desactivado en perfil."

    if forzar_envio:
        return True, None

    now = datetime.now()
    quiet = _quiet_days_set(prefs["quiet_days"])
    if now.weekday() in quiet:
        return False, "Silencio activo para este dia."

    franja = "manana" if now.hour < 15 else "noche"
    expected = prefs["morning_hour"] if franja == "manana" else prefs["night_hour"]
    if abs(now.hour - int(expected)) > 2:
        return False, f"Fuera de ventana {franja} configurada ({expected}:00)."

    return True, None


def _formato_tiempo_restante(segundos):
    if segundos <= 60:
        return f"{segundos} segundos"
    minutos = (segundos + 59) // 60
    return f"{minutos} minutos"


def _rate_limit_config_for_action(accion):
    config = RATE_LIMIT_CONFIG.get(accion)
    if config:
        return config
    if str(accion or "").startswith("login_user:"):
        return RATE_LIMIT_CONFIG.get("login")
    return None


def _accion_rate_login_usuario(username):
    limpio = re.sub(r"[^a-z0-9_.-]+", "", (username or "").strip().lower())
    if not limpio:
        return ""
    return f"login_user:{limpio[:40]}"


def verificar_rate_limit(accion, ip):
    config = _rate_limit_config_for_action(accion)
    if not config:
        return False, 0

    ahora = _segundos_actuales()
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT intentos, ventana_inicio, bloqueado_hasta FROM rate_limits WHERE accion = ? AND ip = ?",
        (accion, ip),
    )
    fila = cursor.fetchone()

    if fila is None:
        conn.close()
        return False, 0

    bloqueado_hasta = fila["bloqueado_hasta"] or 0
    if bloqueado_hasta > ahora:
        conn.close()
        return True, bloqueado_hasta - ahora

    if ahora - fila["ventana_inicio"] >= config["window_seconds"]:
        cursor.execute(
            "UPDATE rate_limits SET intentos = 0, ventana_inicio = ?, bloqueado_hasta = NULL WHERE accion = ? AND ip = ?",
            (ahora, accion, ip),
        )
        conn.commit()

    conn.close()
    return False, 0


def registrar_intento_rate_limit(accion, ip):
    config = _rate_limit_config_for_action(accion)
    if not config:
        return None

    ahora = _segundos_actuales()
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT intentos, ventana_inicio, bloqueado_hasta FROM rate_limits WHERE accion = ? AND ip = ?",
        (accion, ip),
    )
    fila = cursor.fetchone()

    if fila is None:
        cursor.execute(
            """
            INSERT INTO rate_limits (accion, ip, intentos, ventana_inicio, bloqueado_hasta)
            VALUES (?, ?, 1, ?, NULL)
            """,
            (accion, ip, ahora),
        )
        conn.commit()
        conn.close()
        return {"intentos": 1, "bloqueado_hasta": None}

    if ahora - fila["ventana_inicio"] >= config["window_seconds"]:
        intentos = 1
        ventana_inicio = ahora
    else:
        intentos = fila["intentos"] + 1
        ventana_inicio = fila["ventana_inicio"]

    bloqueado_hasta = None
    if intentos >= config["max_attempts"]:
        bloqueado_hasta = ahora + config["block_seconds"]

    cursor.execute(
        """
        UPDATE rate_limits
        SET intentos = ?, ventana_inicio = ?, bloqueado_hasta = ?
        WHERE accion = ? AND ip = ?
        """,
        (intentos, ventana_inicio, bloqueado_hasta, accion, ip),
    )
    conn.commit()
    conn.close()
    return {"intentos": intentos, "bloqueado_hasta": bloqueado_hasta}


def limpiar_rate_limit(accion, ip):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM rate_limits WHERE accion = ? AND ip = ?", (accion, ip))
    conn.commit()
    conn.close()


def registrar_login_attempt(username, ip, success):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO auth_login_attempts (username, ip, success, creado_en)
        VALUES (?, ?, ?, ?)
        """,
        (
            (username or "").strip().lower()[:80],
            (ip or "").strip()[:64],
            1 if success else 0,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    )
    conn.commit()
    conn.close()


def aplicar_rate_limit_api(accion):
    ip_cliente = obtener_ip_cliente()
    bloqueado, segundos_restantes = verificar_rate_limit(accion, ip_cliente)
    if bloqueado:
        return (
            jsonify(
                {
                    "ok": False,
                    "error": "rate_limited",
                    "detail": f"Intenta de nuevo en {_formato_tiempo_restante(segundos_restantes)}.",
                }
            ),
            429,
        )
    registrar_intento_rate_limit(accion, ip_cliente)
    return None


def facebook_configurado():
    return bool(FACEBOOK_APP_ID and FACEBOOK_APP_SECRET)


def google_configurado():
    return bool(GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET)


def _oauth_redirect_uri(endpoint, explicit_uri=""):
    uri = (explicit_uri or "").strip()
    if uri:
        return uri
    if has_request_context():
        scheme = (request.headers.get("X-Forwarded-Proto") or request.scheme or "https").split(",")[0].strip()
        host = (request.headers.get("X-Forwarded-Host") or request.host or "").split(",")[0].strip()
        path = url_for(endpoint)
        if host:
            return f"{scheme}://{host}{path}"
    return url_for(endpoint, _external=True)


def microsoft_configurado():
    return bool(MICROSOFT_CLIENT_ID and MICROSOFT_CLIENT_SECRET)


def whatsapp_configurado():
    return bool(WHATSAPP_TOKEN and WHATSAPP_PHONE_NUMBER_ID)


def sms_configurado():
    return bool(TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN and TWILIO_FROM_NUMBER)


def _mascarar_telefono(telefono):
    raw = "".join(ch for ch in str(telefono or "") if ch.isdigit())
    if len(raw) < 6:
        return str(telefono or "")
    pref = raw[:5]
    suf = raw[-2:]
    return f"+{pref}{'X' * max(0, len(raw) - 7)}{suf}"


def _ultimo_sms_detalle_logs():
    ruta = os.path.join(BASE_DIR, "logs", "notificaciones_last.json")
    if not os.path.exists(ruta):
        return ""
    try:
        with open(ruta, "r", encoding="utf-8") as fh:
            data = json.load(fh) or {}
        sms = data.get("sms") or {}
        return str(sms.get("detalle") or "")
    except Exception:  # noqa: BLE001
        return ""


def obtener_estado_sms_operativo():
    estado = {
        "key": "not_configured",
        "label": "SMS no configurado",
        "detalle": "Faltan variables TWILIO en el servidor.",
        "destinos": [],
        "trial_blocked": False,
    }
    if not sms_configurado():
        return estado

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT telefono FROM usuarios WHERE telefono IS NOT NULL AND trim(telefono) <> ''")
    destinos = [normalizar_telefono_sms(row["telefono"]) for row in cursor.fetchall()]
    conn.close()
    destinos = [d for d in destinos if d]

    detalle_sms = _ultimo_sms_detalle_logs()
    trial_blocked = "21608" in detalle_sms or "unverified" in detalle_sms.lower()
    if trial_blocked:
        return {
            "key": "trial_blocked",
            "label": "SMS bloqueado por Twilio Trial",
            "detalle": "Hay numeros no verificados para cuenta Trial.",
            "destinos": [_mascarar_telefono(d) for d in destinos],
            "trial_blocked": True,
        }

    return {
        "key": "ready",
        "label": "SMS listo",
        "detalle": "Canal SMS configurado y sin bloqueo Trial detectado en el ultimo envio.",
        "destinos": [_mascarar_telefono(d) for d in destinos],
        "trial_blocked": False,
    }


def push_web_configurado():
    return bool(
        PUSH_VAPID_PRIVATE_KEY
        and push_api.push_disponible()
        and _vapid_public_key_valida(PUSH_VAPID_PUBLIC_KEY)
    )


def push_public_key():
    return PUSH_VAPID_PUBLIC_KEY


def push_web_estado():
    return {
        "push_disponible": push_api.push_disponible(),
        "public_key_valida": _vapid_public_key_valida(PUSH_VAPID_PUBLIC_KEY),
        "private_key_present": bool(PUSH_VAPID_PRIVATE_KEY),
        "public_len": len(PUSH_VAPID_PUBLIC_KEY or ""),
        "private_len": len(PUSH_VAPID_PRIVATE_KEY or ""),
        "python_executable": sys.executable,
    }


def _vapid_public_key_valida(value):
    if not value:
        return False
    key = value.strip()
    if "BEGIN PUBLIC KEY" in key or "END PUBLIC KEY" in key:
        return False
    try:
        padded = key + "=" * ((4 - (len(key) % 4)) % 4)
        raw = base64.urlsafe_b64decode(padded.encode("utf-8"))
    except Exception:  # noqa: BLE001
        return False
    return len(raw) == 65 and raw[0] == 4


def guardar_suscripcion_push(usuario_id, subscription):
    endpoint = str(subscription.get("endpoint", "")).strip()
    keys = subscription.get("keys") or {}
    p256dh = str(keys.get("p256dh", "")).strip()
    auth = str(keys.get("auth", "")).strip()
    if not endpoint or not p256dh or not auth:
        raise ValueError("Suscripcion push invalida.")

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO web_push_subscriptions (user_id, endpoint, p256dh, auth, creado_en, actualizado_en)
        VALUES (?, ?, ?, ?, datetime('now'), datetime('now'))
        ON CONFLICT(endpoint) DO UPDATE SET
            user_id = excluded.user_id,
            p256dh = excluded.p256dh,
            auth = excluded.auth,
            actualizado_en = datetime('now')
        """,
        (usuario_id, endpoint, p256dh, auth),
    )
    conn.commit()
    conn.close()


def obtener_suscripciones_push(usuario_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT endpoint, p256dh, auth FROM web_push_subscriptions WHERE user_id = ?",
        (usuario_id,),
    )
    filas = cursor.fetchall()
    conn.close()
    return [
        {
            "endpoint": fila["endpoint"],
            "keys": {"p256dh": fila["p256dh"], "auth": fila["auth"]},
        }
        for fila in filas
    ]


def eliminar_suscripcion_push(endpoint):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM web_push_subscriptions WHERE endpoint = ?", (endpoint,))
    conn.commit()
    conn.close()


def enviar_push_usuario(usuario_id, title, body, url="/notificaciones"):
    if not push_web_configurado():
        raise RuntimeError("Push web no configurado en el servidor.")

    suscripciones = obtener_suscripciones_push(usuario_id)
    if not suscripciones:
        return {"enviados": 0, "expiradas": 0}

    payload = {"title": title, "body": body, "url": url}
    enviados = 0
    expiradas = 0

    for sub in suscripciones:
        try:
            push_api.enviar_web_push(
                sub,
                PUSH_VAPID_PRIVATE_KEY,
                {"sub": PUSH_VAPID_SUBJECT},
                payload,
            )
            enviados += 1
        except Exception as exc:  # noqa: BLE001
            status_code = getattr(getattr(exc, "response", None), "status_code", None)
            if status_code in (404, 410):
                eliminar_suscripcion_push(sub["endpoint"])
                expiradas += 1
            else:
                app.logger.warning("No se pudo enviar push web: %s", exc)

    return {"enviados": enviados, "expiradas": expiradas}


def construir_mensaje_push_resumen(notis):
    return notificaciones_compose.construir_push_resumen(notis, FORCE_NOTIFICATIONS_TEST)


def enviar_recordatorio_push_usuario(usuario_id, forzar_envio=False):
    if not push_web_configurado():
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": "Push web no configurado."}
    ok_canal, detalle_canal = _canal_notificacion_habilitado(usuario_id, "push", forzar_envio)
    if not ok_canal:
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": detalle_canal}

    suscripciones = obtener_suscripciones_push(usuario_id)
    if not suscripciones:
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Sin suscripciones push."}

    notis = obtener_notificaciones_usuario(usuario_id)
    mensaje = construir_mensaje_push_resumen(notis)
    if not mensaje:
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Sin alertas para enviar."}

    marca_envio = marca_envio_actual()
    conn = get_connection()
    cursor = conn.cursor()
    if not forzar_envio:
        cursor.execute(
            "SELECT 1 FROM web_push_notificaciones_log WHERE user_id = ? AND fecha_envio = ?",
            (usuario_id, marca_envio),
        )
        if cursor.fetchone():
            conn.close()
            return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Push ya enviado en esta franja."}

    try:
        resultado = enviar_push_usuario(usuario_id, "Meta Inteligente", mensaje, "/notificaciones")
        if resultado["enviados"] > 0:
            cursor.execute(
                """
                INSERT OR IGNORE INTO web_push_notificaciones_log (user_id, fecha_envio, creado_en)
                VALUES (?, ?, ?)
                """,
                (usuario_id, marca_envio, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            )
            conn.commit()
        conn.close()
        return {
            "enviados": resultado["enviados"],
            "omitidos": 0 if resultado["enviados"] > 0 else 1,
            "errores": 0,
            "detalle": None if resultado["enviados"] > 0 else "No hay suscripciones activas.",
        }
    except Exception as exc:  # noqa: BLE001
        conn.close()
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": str(exc)}


def limpiar_telefono(telefono):
    limpio = re.sub(r"[^0-9]", "", (telefono or ""))
    if not limpio:
        return ""
    return limpio


def normalizar_telefono_sms(telefono):
    digitos = limpiar_telefono(telefono)
    if not digitos:
        return ""

    prefijo = SMS_DEFAULT_COUNTRY_CODE
    if not prefijo.startswith("+"):
        prefijo = f"+{prefijo}"
    prefijo_digitos = re.sub(r"[^0-9]", "", prefijo)

    if telefono and str(telefono).strip().startswith("+"):
        return f"+{digitos}"
    # Si ya viene en formato internacional sin "+", no anteponer prefijo por defecto.
    if len(digitos) >= 11:
        return f"+{digitos}"
    if prefijo_digitos and digitos.startswith(prefijo_digitos):
        return f"+{digitos}"
    return f"{prefijo}{digitos}"


def extraer_message_ids_whatsapp(response_json):
    if not isinstance(response_json, dict):
        return []
    mensajes = response_json.get("messages")
    if not isinstance(mensajes, list):
        return []
    ids = []
    for item in mensajes:
        if isinstance(item, dict) and item.get("id"):
            ids.append(str(item["id"]))
    return ids


def registrar_log_whatsapp_envio(user_id, meta_id, tipo, telefono, message_ids):
    logs_dir = os.path.join(BASE_DIR, "logs")
    os.makedirs(logs_dir, exist_ok=True)
    ruta = os.path.join(logs_dir, "whatsapp_message_ids.log")
    marca = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ids_texto = ",".join(message_ids) if message_ids else "-"
    linea = (
        f"[{marca}] user_id={user_id} meta_id={meta_id} tipo={tipo} "
        f"telefono={telefono} message_ids={ids_texto}\n"
    )
    with open(ruta, "a", encoding="utf-8") as f:
        f.write(linea)


def username_slug(texto):
    base = re.sub(r"[^a-z0-9_]+", "_", (texto or "").strip().lower())
    base = base.strip("_")
    return base or "usuario"


def generar_username_disponible(base_username, cursor):
    base = username_slug(base_username)
    candidato = base
    idx = 1
    while True:
        cursor.execute("SELECT 1 FROM usuarios WHERE username = ?", (candidato,))
        if cursor.fetchone() is None:
            return candidato
        idx += 1
        candidato = f"{base}_{idx}"


def crear_o_obtener_usuario_facebook(fb_id, fb_name, fb_email):
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT id, username FROM usuarios WHERE facebook_id = ?", (fb_id,))
    user = cursor.fetchone()
    if user:
        conn.close()
        return user["id"], user["username"]

    email = (fb_email or "").strip().lower() or None
    if email:
        cursor.execute("SELECT id, username FROM usuarios WHERE email = ?", (email,))
        existente = cursor.fetchone()
        if existente:
            cursor.execute("UPDATE usuarios SET facebook_id = ? WHERE id = ?", (fb_id, existente["id"]))
            conn.commit()
            conn.close()
            return existente["id"], existente["username"]

    username_base = fb_name or f"fb_{fb_id[-6:]}"
    username = generar_username_disponible(username_base, cursor)
    password_temp = hash_contrasena(secrets.token_urlsafe(16))
    cursor.execute(
        "INSERT INTO usuarios (username, password, email, facebook_id, rol) VALUES (?, ?, ?, ?, ?)",
        (username, password_temp, email, fb_id, "editor"),
    )
    conn.commit()
    user_id = cursor.lastrowid
    conn.close()
    return user_id, username


def crear_o_obtener_usuario_oauth(provider, provider_id, nombre, email):
    conn = get_connection()
    cursor = conn.cursor()

    email_normalizado = (email or "").strip().lower() or None
    if email_normalizado:
        cursor.execute("SELECT id, username FROM usuarios WHERE lower(email) = ?", (email_normalizado,))
        existente = cursor.fetchone()
        if existente:
            conn.close()
            return existente["id"], existente["username"]

    username_base = email_normalizado.split("@")[0] if email_normalizado else (nombre or f"{provider}_{provider_id[-6:]}")
    username = generar_username_disponible(username_base, cursor)
    password_temp = hash_contrasena(secrets.token_urlsafe(16))
    cursor.execute(
        "INSERT INTO usuarios (username, password, email, rol) VALUES (?, ?, ?, ?)",
        (username, password_temp, email_normalizado, "editor"),
    )
    conn.commit()
    user_id = cursor.lastrowid
    conn.close()
    return user_id, username


def enviar_whatsapp_codigo(destino_telefono, codigo):
    if not whatsapp_configurado():
        raise RuntimeError("WhatsApp no configurado.")
    return whatsapp_api.send_text(
        WHATSAPP_API_VERSION,
        WHATSAPP_PHONE_NUMBER_ID,
        WHATSAPP_TOKEN,
        destino_telefono,
        f"Tu codigo de verificacion Meta90 es: {codigo}. Expira en 10 minutos.",
    )


def enviar_whatsapp_texto(destino_telefono, texto):
    if not whatsapp_configurado():
        raise RuntimeError("WhatsApp no configurado.")
    return whatsapp_api.send_text(
        WHATSAPP_API_VERSION,
        WHATSAPP_PHONE_NUMBER_ID,
        WHATSAPP_TOKEN,
        destino_telefono,
        texto,
    )


def enviar_whatsapp_template(destino_telefono, template_name=None, template_lang=None):
    if not whatsapp_configurado():
        raise RuntimeError("WhatsApp no configurado.")

    nombre = (template_name or WHATSAPP_TEMPLATE_NAME).strip()
    idioma = (template_lang or WHATSAPP_TEMPLATE_LANG).strip()
    if not nombre:
        raise RuntimeError("No hay plantilla de WhatsApp configurada.")

    return whatsapp_api.send_template(
        WHATSAPP_API_VERSION,
        WHATSAPP_PHONE_NUMBER_ID,
        WHATSAPP_TOKEN,
        destino_telefono,
        nombre,
        idioma,
    )


def enviar_whatsapp_notificacion(destino_telefono, texto):
    modo = WHATSAPP_NOTI_MODE if WHATSAPP_NOTI_MODE in (
        "template_only",
        "text_only",
        "template_then_text",
    ) else "template_then_text"

    if modo == "text_only":
        try:
            resp = enviar_whatsapp_texto(destino_telefono, texto)
            return extraer_message_ids_whatsapp(resp)
        except Exception as exc:  # noqa: BLE001
            # En sandbox, texto libre puede fallar fuera de ventana de 24h.
            app.logger.warning("Texto directo bloqueado, intentando plantilla: %s", exc)
            resp = enviar_whatsapp_template(destino_telefono)
            return extraer_message_ids_whatsapp(resp)

    if modo == "template_only":
        resp = enviar_whatsapp_template(destino_telefono)
        return extraer_message_ids_whatsapp(resp)

    # Modo recomendado para sandbox: abre conversacion con plantilla y luego intenta texto.
    ids = []
    resp_tpl = enviar_whatsapp_template(destino_telefono)
    ids.extend(extraer_message_ids_whatsapp(resp_tpl))
    try:
        resp_txt = enviar_whatsapp_texto(destino_telefono, texto)
        ids.extend(extraer_message_ids_whatsapp(resp_txt))
    except Exception as exc:  # noqa: BLE001
        app.logger.warning("Texto adicional de WhatsApp no entregado: %s", exc)
    return ids


def enviar_sms_notificacion(destino_telefono, texto):
    if not sms_configurado():
        raise RuntimeError("SMS no configurado.")
    numero = normalizar_telefono_sms(destino_telefono)
    if not numero:
        raise RuntimeError("Telefono invalido para SMS.")
    return sms_api.send_twilio_sms(
        TWILIO_ACCOUNT_SID,
        TWILIO_AUTH_TOKEN,
        TWILIO_FROM_NUMBER,
        numero,
        texto,
    )


def construir_mensajes_notificacion_movil(notis):
    return notificaciones_compose.construir_sms_items(notis)


def expandir_mensajes_sms_prueba(mensajes):
    if SMS_TEST_MIN_NOTIFICATIONS <= 0 or len(mensajes) >= SMS_TEST_MIN_NOTIFICATIONS:
        return mensajes
    if not mensajes:
        return mensajes

    # En modo prueba, duplicamos mensajes existentes para alcanzar el minimo solicitado.
    extendidos = list(mensajes)
    i = 0
    while len(extendidos) < SMS_TEST_MIN_NOTIFICATIONS:
        base = mensajes[i % len(mensajes)]
        extendidos.append(
            {
                "meta_id": base["meta_id"],
                "tipo": f"{base['tipo']}_test_{len(extendidos)+1}",
                "texto": f"[PRUEBA {len(extendidos)+1}] {base['texto']}",
            }
        )
        i += 1
    return extendidos


def enviar_recordatorios_sms_usuario(usuario_id, forzar_envio=False):
    ok_canal, detalle_canal = _canal_notificacion_habilitado(usuario_id, "sms", forzar_envio)
    if not ok_canal:
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": detalle_canal}

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT username, telefono FROM usuarios WHERE id = ?", (usuario_id,))
    user = cursor.fetchone()
    if user is None:
        conn.close()
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": "Usuario no encontrado."}

    telefono = normalizar_telefono_sms(user["telefono"])
    if not telefono:
        conn.close()
        return {
            "enviados": 0,
            "omitidos": 0,
            "errores": 1,
            "detalle": "No tienes telefono configurado en tu perfil.",
        }

    notis = obtener_notificaciones_usuario(usuario_id)
    pendientes = construir_mensajes_notificacion_movil(notis)
    if not pendientes and FORCE_NOTIFICATIONS_TEST:
        pendientes = [
            {
                "meta_id": 0,
                "tipo": "test",
                "texto": "Mensaje de prueba: canal SMS activo.",
            }
        ]
    pendientes = expandir_mensajes_sms_prueba(pendientes)
    marca_envio = marca_envio_actual()

    enviados = 0
    omitidos = 0
    errores = 0
    ultimo_error = None

    for item in pendientes:
        if not forzar_envio:
            cursor.execute(
                """
                SELECT 1 FROM sms_notificaciones_log
                WHERE user_id = ? AND meta_id = ? AND tipo = ? AND fecha_envio = ?
                """,
                (usuario_id, item["meta_id"], item["tipo"], marca_envio),
            )
            if cursor.fetchone():
                omitidos += 1
                continue

        try:
            enviar_sms_notificacion(telefono, item["texto"])
            cursor.execute(
                """
                INSERT OR IGNORE INTO sms_notificaciones_log (user_id, meta_id, tipo, fecha_envio, creado_en)
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    usuario_id,
                    item["meta_id"],
                    item["tipo"],
                    marca_envio,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                ),
            )
            conn.commit()
            enviados += 1
        except Exception as exc:  # noqa: BLE001
            errores += 1
            ultimo_error = str(exc)
            if "63038" in ultimo_error:
                # Limite diario de Twilio alcanzado; evitar reintentos inutiles.
                break

    conn.close()
    return {
        "enviados": enviados,
        "omitidos": omitidos,
        "errores": errores,
        "detalle": ultimo_error,
    }


def enviar_recordatorios_movil_usuario(usuario_id, forzar_envio=False):
    return enviar_recordatorios_sms_usuario(usuario_id, forzar_envio)


def enviar_recordatorios_push_todos(forzar_envio=False):
    if not push_web_configurado():
        return {
            "usuarios_total": 0,
            "enviados": 0,
            "omitidos": 0,
            "errores": 0,
            "detalle": "Push web no configurado.",
        }

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT user_id FROM web_push_subscriptions")
    usuarios = [fila["user_id"] for fila in cursor.fetchall()]
    conn.close()

    resumen = {
        "usuarios_total": len(usuarios),
        "enviados": 0,
        "omitidos": 0,
        "errores": 0,
        "detalle": None,
    }

    for usuario_id in usuarios:
        resultado = enviar_recordatorio_push_usuario(usuario_id, forzar_envio)
        resumen["enviados"] += resultado["enviados"]
        resumen["omitidos"] += resultado["omitidos"]
        resumen["errores"] += resultado["errores"]
        if resultado["detalle"]:
            resumen["detalle"] = resultado["detalle"]

    return resumen


def construir_email_notificaciones(username, notis):
    return notificaciones_compose.construir_email(username, notis, FORCE_NOTIFICATIONS_TEST)


def enviar_recordatorios_sms_todos(forzar_envio=False):
    if notificaciones_en_mantenimiento():
        return {
            "usuarios_total": 0,
            "enviados": 0,
            "omitidos": 0,
            "errores": 0,
            "sms_usuarios_total": 0,
            "sms_enviados": 0,
            "sms_omitidos": 0,
            "sms_errores": 0,
            "sms_detalle": "Notificaciones en mantenimiento.",
            "push_usuarios_total": 0,
            "push_enviados": 0,
            "push_omitidos": 0,
            "push_errores": 0,
            "push_detalle": "Notificaciones en mantenimiento.",
            "email_usuarios_total": 0,
            "email_enviados": 0,
            "email_omitidos": 0,
            "email_errores": 0,
            "email_detalle": "Notificaciones en mantenimiento.",
        }

    usuarios = []
    sms_detalle = None
    if sms_configurado():
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM usuarios WHERE telefono IS NOT NULL AND trim(telefono) <> ''")
        usuarios = [fila["id"] for fila in cursor.fetchall()]
        conn.close()
    else:
        sms_detalle = "SMS no configurado."

    resumen = {
        "usuarios_total": len(usuarios),
        "enviados": 0,
        "omitidos": 0,
        "errores": 0,
    }

    for usuario_id in usuarios:
        resultado = enviar_recordatorios_sms_usuario(usuario_id, forzar_envio)
        resumen["enviados"] += resultado["enviados"]
        resumen["omitidos"] += resultado["omitidos"]
        resumen["errores"] += resultado["errores"]
        if resultado.get("detalle"):
            sms_detalle = resultado["detalle"]

    push_resumen = enviar_recordatorios_push_todos(forzar_envio)
    resumen["push_usuarios_total"] = push_resumen["usuarios_total"]
    resumen["push_enviados"] = push_resumen["enviados"]
    resumen["push_omitidos"] = push_resumen["omitidos"]
    resumen["push_errores"] = push_resumen["errores"]
    if push_resumen.get("detalle"):
        resumen["push_detalle"] = push_resumen["detalle"]

    email_resumen = enviar_recordatorios_email_todos(forzar_envio)
    resumen["email_usuarios_total"] = email_resumen["usuarios_total"]
    resumen["email_enviados"] = email_resumen["enviados"]
    resumen["email_omitidos"] = email_resumen["omitidos"]
    resumen["email_errores"] = email_resumen["errores"]
    if email_resumen.get("detalle"):
        resumen["email_detalle"] = email_resumen["detalle"]

    resumen["sms_usuarios_total"] = resumen["usuarios_total"]
    resumen["sms_enviados"] = resumen["enviados"]
    resumen["sms_omitidos"] = resumen["omitidos"]
    resumen["sms_errores"] = resumen["errores"]
    if sms_detalle:
        resumen["sms_detalle"] = sms_detalle

    return resumen


def enviar_recordatorios_movil_todos():
    return enviar_recordatorios_sms_todos()


def iniciar_sesion_usuario(user_id, username, rol=None):
    rol_limpio = (rol or "").strip().lower()
    if rol_limpio not in ("admin", "editor", "lector"):
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT rol FROM usuarios WHERE id = ?", (user_id,))
        fila = cursor.fetchone()
        conn.close()
        rol_limpio = (fila["rol"] if fila and fila["rol"] else "editor").strip().lower() if fila else "editor"
        if rol_limpio not in ("admin", "editor", "lector"):
            rol_limpio = "editor"

    session["usuario"] = username
    session["usuario_id"] = user_id
    session["usuario_rol"] = rol_limpio
    session["csrf_token"] = secrets.token_urlsafe(32)


def _admin_2fa_requerido_para_user(user):
    if app.config.get("TESTING"):
        return False
    if not ADMIN_2FA_REQUIRED:
        return False
    if not user:
        return False
    rol = str(user.get("rol") if isinstance(user, dict) else user["rol"]).strip().lower()
    return rol == "admin" and proveedor_correo_configurado()


def iniciar_admin_2fa_pendiente(user):
    code = f"{secrets.randbelow(1000000):06d}"
    expira = (datetime.now() + timedelta(minutes=10)).strftime("%Y-%m-%d %H:%M:%S")
    payload = {"id": int(user["id"]), "username": str(user["username"]), "rol": str(user["rol"]), "code": code, "expira": expira}
    session["admin_2fa_pending"] = payload
    asunto = "Codigo 2FA Admin - Meta90"
    cuerpo = (
        "Se detecto inicio de sesion admin.\n\n"
        f"Tu codigo 2FA es: {code}\n"
        "Expira en 10 minutos.\n\n"
        "Si no fuiste tu, cambia tu clave."
    )
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT email FROM usuarios WHERE id = ?", (int(user["id"]),))
    row = cursor.fetchone()
    conn.close()
    destino = (row["email"] if row and row["email"] else "").strip().lower()
    if email_valido(destino):
        enviar_email_generico(destino, asunto, cuerpo)
        return True, None
    return False, "Admin sin email valido para 2FA."


def validar_admin_2fa_pending(codigo):
    payload = session.get("admin_2fa_pending")
    if not payload:
        return False, "No hay verificacion pendiente."
    exp = parsear_fecha_hora(payload.get("expira", ""))
    if exp is None or exp < datetime.now():
        session.pop("admin_2fa_pending", None)
        return False, "El codigo expiro."
    if str(codigo or "").strip() != str(payload.get("code", "")):
        return False, "Codigo incorrecto."
    return True, payload


def _serializer_recordar():
    return URLSafeSerializer(app.secret_key, salt="meta90-remember-v1")


def generar_token_recordar(user_id, username):
    return _serializer_recordar().dumps({"uid": int(user_id), "usr": str(username or "")})


def validar_token_recordar(token):
    try:
        payload = _serializer_recordar().loads(token)
    except BadSignature:
        return None

    user_id = int(payload.get("uid", 0))
    username = str(payload.get("usr", "")).strip()
    if user_id <= 0 or not username:
        return None

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, rol FROM usuarios WHERE id = ?", (user_id,))
    user = cursor.fetchone()
    conn.close()
    if user is None:
        return None
    if str(user["username"]).strip() != username:
        return None
    return {"id": user["id"], "username": user["username"], "rol": user["rol"]}


def _hash_api_token(token):
    return hashlib.sha256(str(token or "").encode("utf-8")).hexdigest()


def generar_api_token_usuario(user_id, name="principal"):
    raw = "mt90_" + secrets.token_urlsafe(32)
    token_hash = _hash_api_token(raw)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO api_tokens (user_id, name, token_hash, active, created_at)
        VALUES (?, ?, ?, 1, ?)
        """,
        (user_id, (name or "principal")[:50], token_hash, now),
    )
    conn.commit()
    conn.close()
    return raw


def listar_api_tokens_usuario(user_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, name, active, created_at, last_used_at
        FROM api_tokens
        WHERE user_id = ?
        ORDER BY created_at DESC
        """,
        (user_id,),
    )
    rows = cursor.fetchall()
    conn.close()
    return rows


def revocar_api_token_usuario(user_id, token_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE api_tokens SET active = 0 WHERE id = ? AND user_id = ?",
        (token_id, user_id),
    )
    conn.commit()
    ok = cursor.rowcount > 0
    conn.close()
    return ok


def autenticar_api_token(header_value):
    raw = (header_value or "").strip()
    if raw.lower().startswith("bearer "):
        raw = raw[7:].strip()
    if not raw:
        return None

    token_hash = _hash_api_token(raw)
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT t.id, t.user_id, u.username, coalesce(u.rol, 'editor') AS rol
        FROM api_tokens t
        JOIN usuarios u ON u.id = t.user_id
        WHERE t.token_hash = ? AND t.active = 1
        """,
        (token_hash,),
    )
    row = cursor.fetchone()
    if row:
        cursor.execute(
            "UPDATE api_tokens SET last_used_at = ? WHERE id = ?",
            (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), row["id"]),
        )
        conn.commit()
    conn.close()
    if not row:
        return None
    return {"user_id": row["user_id"], "username": row["username"], "rol": row["rol"]}


@app.before_request
def restaurar_sesion_recordada():
    if usuario_autenticado():
        return
    if request.endpoint == "static":
        return

    token = (request.cookies.get("recordar_token") or "").strip()
    if not token:
        return

    remembered_session = validar_token_recordar(token)
    if not remembered_session:
        return

    iniciar_sesion_usuario(
        remembered_session["id"],
        remembered_session["username"],
        remembered_session.get("rol"),
    )
    session.permanent = True


@app.before_request
def validar_csrf():
    if app.config.get("TESTING"):
        return
    if request.method in ("GET", "HEAD", "OPTIONS", "TRACE"):
        return
    if request.endpoint == "static":
        return
    if not usuario_autenticado():
        return

    token_session = session.get("csrf_token")
    if not token_session:
        abort(400, description="CSRF token faltante.")

    token_enviado = (
        request.headers.get("X-CSRF-Token", "").strip()
        or request.form.get("_csrf_token", "").strip()
    )
    if not token_enviado or not secrets.compare_digest(token_session, token_enviado):
        abort(400, description="CSRF token invalido.")


@app.before_request
def validar_permiso_modulo():
    if request.endpoint == "static":
        return
    if not usuario_autenticado():
        return
    if usuario_es_admin():
        return

    modulo = _modulo_from_endpoint(request.endpoint)
    if not modulo:
        return

    user_id = obtener_usuario_id_actual()
    permiso = obtener_permiso_modulo_usuario(user_id, modulo)
    is_write = request.method not in ("GET", "HEAD", "OPTIONS")
    allowed = permiso["edit"] if is_write else permiso["view"]
    if allowed:
        return

    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": "forbidden", "detail": "Permiso insuficiente."}), 403
    flash("No tienes permiso para acceder a este modulo.", "warning")
    return redirect(url_for("metas_bp.inicio_metas"))


@app.after_request
def aplicar_headers_seguridad(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    if request.is_secure:
        response.headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    return response


def smtp_configurado():
    return bool(SMTP_HOST and SMTP_PORT and SMTP_USER and SMTP_PASSWORD and SMTP_FROM)


def resend_configurado():
    if not (RESEND_API_KEY and RESEND_FROM):
        return False
    # Evita tomar placeholders como claves validas.
    return "xxxxxxxx" not in RESEND_API_KEY.lower() and "pega_aqui" not in RESEND_API_KEY.lower()


def proveedor_correo_configurado():
    return resend_configurado() or smtp_configurado()


def ruta_proveedores_correo():
    proveedor = MAIL_PROVIDER if MAIL_PROVIDER in ("auto", "resend", "smtp") else "auto"
    if proveedor == "smtp":
        return ["SMTP", "Resend"]
    # auto y resend priorizan Resend con fallback SMTP.
    return ["Resend", "SMTP"]


def nombre_ruta_proveedores():
    disponibles = []
    for nombre in ruta_proveedores_correo():
        if nombre == "Resend" and resend_configurado():
            disponibles.append("Resend")
        if nombre == "SMTP" and smtp_configurado():
            disponibles.append("SMTP")
    return " -> ".join(disponibles) if disponibles else "Sin configurar"


def email_valido(email):
    email = (email or "").strip()
    return "@" in email and "." in email.split("@")[-1]


def enmascarar_email(email):
    email = (email or "").strip()
    if "@" not in email:
        return email
    usuario, dominio = email.split("@", 1)
    if len(usuario) <= 1:
        usuario_mask = "*"
    elif len(usuario) == 2:
        usuario_mask = usuario[0] + "*"
    else:
        usuario_mask = usuario[0] + ("*" * (len(usuario) - 2)) + usuario[-1]
    return f"{usuario_mask}@{dominio}"


def enviar_email_recuperacion(destino_email, codigo_recuperacion):
    asunto = "Recuperacion de cuenta - Meta Inteligente"
    cuerpo = (
        "Recibimos una solicitud para recuperar tu cuenta.\n\n"
        f"Tu codigo de recuperacion es: {codigo_recuperacion}\n\n"
        "El codigo expira en 15 minutos. Si no solicitaste este cambio, ignora este mensaje."
    )

    return enviar_email_generico(destino_email, asunto, cuerpo)


def enviar_email_generico(destino_email, asunto, cuerpo):
    errores = []

    for proveedor in ruta_proveedores_correo():
        if proveedor == "Resend":
            if not resend_configurado():
                continue

            try:
                mail_api.send_resend(RESEND_API_KEY, RESEND_FROM, destino_email, asunto, cuerpo)
                return "Resend"
            except Exception as exc:  # noqa: BLE001
                errores.append(f"Resend error: {exc}")

        if proveedor == "SMTP":
            if not smtp_configurado():
                continue

            try:
                mail_api.send_smtp(
                    SMTP_HOST,
                    SMTP_PORT,
                    SMTP_USER,
                    SMTP_PASSWORD,
                    SMTP_FROM,
                    destino_email,
                    asunto,
                    cuerpo,
                    SMTP_USE_TLS,
                )
                return "SMTP"
            except Exception as exc:  # noqa: BLE001
                errores.append(f"SMTP error: {exc}")

    if errores:
        raise RuntimeError(" | ".join(errores))
    raise RuntimeError("No hay proveedor de correo configurado (Resend o SMTP).")


def enviar_recordatorio_email_usuario(usuario_id, forzar_envio=False):
    if not proveedor_correo_configurado():
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": "Correo no configurado."}
    ok_canal, detalle_canal = _canal_notificacion_habilitado(usuario_id, "email", forzar_envio)
    if not ok_canal:
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": detalle_canal}

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT username, email FROM usuarios WHERE id = ?", (usuario_id,))
    user = cursor.fetchone()
    if user is None:
        conn.close()
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": "Usuario no encontrado."}

    destino_email = (user["email"] or "").strip().lower()
    if not email_valido(destino_email):
        conn.close()
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Usuario sin email valido."}

    notis = obtener_notificaciones_usuario(usuario_id)
    asunto, cuerpo = construir_email_notificaciones(user["username"], notis)
    if not asunto:
        conn.close()
        return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Sin alertas para enviar."}

    marca_envio = marca_envio_actual()
    if not forzar_envio:
        cursor.execute(
            "SELECT 1 FROM email_notificaciones_log WHERE user_id = ? AND fecha_envio = ?",
            (usuario_id, marca_envio),
        )
        if cursor.fetchone():
            conn.close()
            return {"enviados": 0, "omitidos": 1, "errores": 0, "detalle": "Email ya enviado en esta franja."}

    try:
        proveedor = enviar_email_generico(destino_email, asunto, cuerpo)
        cursor.execute(
            """
            INSERT OR IGNORE INTO email_notificaciones_log (user_id, fecha_envio, creado_en)
            VALUES (?, ?, ?)
            """,
            (usuario_id, marca_envio, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        conn.commit()
        conn.close()
        return {"enviados": 1, "omitidos": 0, "errores": 0, "detalle": f"Enviado por {proveedor}."}
    except Exception as exc:  # noqa: BLE001
        conn.close()
        return {"enviados": 0, "omitidos": 0, "errores": 1, "detalle": str(exc)}


def enviar_recordatorios_email_todos(forzar_envio=False):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM usuarios WHERE email IS NOT NULL AND trim(email) <> ''")
    usuarios = [fila["id"] for fila in cursor.fetchall()]
    conn.close()

    resumen = {
        "usuarios_total": len(usuarios),
        "enviados": 0,
        "omitidos": 0,
        "errores": 0,
        "detalle": None,
    }
    for usuario_id in usuarios:
        resultado = enviar_recordatorio_email_usuario(usuario_id, forzar_envio)
        resumen["enviados"] += resultado["enviados"]
        resumen["omitidos"] += resultado["omitidos"]
        resumen["errores"] += resultado["errores"]
        if resultado["detalle"]:
            resumen["detalle"] = resultado["detalle"]
    return resumen


def _descartes_hoy_usuario(usuario_id, fecha_alerta):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT tipo, referencia
        FROM notificaciones_descartadas
        WHERE user_id = ? AND fecha_alerta = ?
        """,
        (usuario_id, fecha_alerta),
    )
    filas = cursor.fetchall()
    conn.close()
    return {(fila["tipo"], fila["referencia"]) for fila in filas}


def descartar_notificacion_usuario(usuario_id, tipo, referencia):
    tipo = (tipo or "").strip().lower()
    referencia = str(referencia or "").strip()
    if not tipo or not referencia:
        return False
    permitidos = {
        "meta_vencida",
        "meta_por_vencer",
        "meta_rezago",
        "calendario_hoy",
        "calendario_proximo",
        "mensaje_no_leido",
    }
    if tipo not in permitidos:
        return False

    fecha_alerta = datetime.now().strftime("%Y-%m-%d")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT OR IGNORE INTO notificaciones_descartadas
        (user_id, tipo, referencia, fecha_alerta, creado_en)
        VALUES (?, ?, ?, ?, ?)
        """,
        (usuario_id, tipo, referencia, fecha_alerta, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    conn.commit()
    inserted = cursor.rowcount > 0
    conn.close()
    return inserted


def descartar_todas_notificaciones_usuario(usuario_id):
    notis = obtener_notificaciones_usuario(usuario_id)
    fecha_alerta = datetime.now().strftime("%Y-%m-%d")
    filas = []

    for item in notis["vencidas"]:
        filas.append(("meta_vencida", str(item["id"])))
    for item in notis["por_vencer"]:
        filas.append(("meta_por_vencer", str(item["id"])))
    for item in notis.get("metas_rezagadas", []):
        filas.append(("meta_rezago", str(item["id"])))
    for item in notis["eventos_hoy"]:
        filas.append(("calendario_hoy", str(item["id"])))
    for item in notis["eventos_proximos"]:
        filas.append(("calendario_proximo", str(item["id"])))
    for item in notis["mensajes_no_leidos_recientes"]:
        filas.append(("mensaje_no_leido", str(item["id"])))

    if not filas:
        return 0

    conn = get_connection()
    cursor = conn.cursor()
    creadas = 0
    marca = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for tipo, referencia in filas:
        cursor.execute(
            """
            INSERT OR IGNORE INTO notificaciones_descartadas
            (user_id, tipo, referencia, fecha_alerta, creado_en)
            VALUES (?, ?, ?, ?, ?)
            """,
            (usuario_id, tipo, referencia, fecha_alerta, marca),
        )
        if cursor.rowcount > 0:
            creadas += 1
    conn.commit()
    conn.close()
    return creadas


def restaurar_notificaciones_descartadas_hoy(usuario_id):
    fecha_alerta = datetime.now().strftime("%Y-%m-%d")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "DELETE FROM notificaciones_descartadas WHERE user_id = ? AND fecha_alerta = ?",
        (usuario_id, fecha_alerta),
    )
    conn.commit()
    eliminadas = cursor.rowcount
    conn.close()
    return eliminadas


def obtener_notificaciones_usuario(usuario_id):
    fecha_alerta = datetime.now().strftime("%Y-%m-%d")
    descartes_hoy = _descartes_hoy_usuario(usuario_id, fecha_alerta)
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, meta, monto, ahorrado, fecha_limite FROM metas WHERE user_id = ?", (usuario_id,))
    metas = cursor.fetchall()
    hoy = datetime.now().date()
    vencidas = []
    por_vencer = []
    metas_rezagadas = []

    for item in metas:
        fecha_obj = parsear_fecha_limite(item["fecha_limite"])
        if not fecha_obj:
            continue

        dias_restantes = (fecha_obj - hoy).days
        payload = {
            "id": item["id"],
            "meta": item["meta"],
            "fecha_limite": fecha_obj.strftime("%Y-%m-%d"),
        }

        if dias_restantes < 0:
            payload["dias"] = abs(dias_restantes)
            if ("meta_vencida", str(payload["id"])) not in descartes_hoy:
                vencidas.append(payload)
        elif dias_restantes <= 3:
            payload["dias"] = dias_restantes
            if ("meta_por_vencer", str(payload["id"])) not in descartes_hoy:
                por_vencer.append(payload)

        monto = float(item["monto"] or 0)
        ahorrado = float(item["ahorrado"] or 0)
        if monto > 0 and dias_restantes >= 0:
            porcentaje = round((ahorrado / monto) * 100.0, 1)
            umbral = 0
            if dias_restantes <= 7:
                umbral = 95
            elif dias_restantes <= 15:
                umbral = 85
            elif dias_restantes <= 30:
                umbral = 70
            if umbral and porcentaje < umbral:
                if ("meta_rezago", str(payload["id"])) not in descartes_hoy:
                    metas_rezagadas.append(
                        {
                            **payload,
                            "dias": dias_restantes,
                            "porcentaje": porcentaje,
                            "faltante": round(max(0.0, monto - ahorrado), 2),
                        }
                    )

    vencidas.sort(key=lambda x: x["dias"], reverse=True)
    por_vencer.sort(key=lambda x: x["dias"])
    metas_rezagadas.sort(key=lambda x: (x["dias"], x["porcentaje"]))
    fecha_hoy = hoy.strftime("%Y-%m-%d")
    fecha_limite = (hoy + timedelta(days=3)).strftime("%Y-%m-%d")
    cursor.execute(
        """
        SELECT id, titulo, fecha_evento, hora_evento
        FROM calendario_eventos
        WHERE user_id = ? AND fecha_evento BETWEEN ? AND ?
        ORDER BY fecha_evento ASC, coalesce(hora_evento, '99:99') ASC, id ASC
        """,
        (usuario_id, fecha_hoy, fecha_limite),
    )
    eventos = cursor.fetchall()

    eventos_hoy = []
    eventos_proximos = []
    for item in eventos:
        fecha_evento = parsear_fecha_limite(item["fecha_evento"])
        if fecha_evento is None:
            continue
        dias_restantes = (fecha_evento - hoy).days
        payload = {
            "id": item["id"],
            "titulo": item["titulo"],
            "fecha_evento": fecha_evento.strftime("%Y-%m-%d"),
            "hora_evento": (item["hora_evento"] or "").strip(),
            "dias": dias_restantes,
        }
        if dias_restantes == 0:
            if ("calendario_hoy", str(payload["id"])) not in descartes_hoy:
                eventos_hoy.append(payload)
        elif dias_restantes > 0:
            if ("calendario_proximo", str(payload["id"])) not in descartes_hoy:
                eventos_proximos.append(payload)

    cursor.execute(
        """
        SELECT id, coalesce(asunto, '(Sin asunto)') AS asunto, coalesce(remitente, '-') AS remitente
        FROM mensajes
        WHERE user_id = ? AND carpeta = 'entrada' AND leido = 0
        ORDER BY datetime(creado_en) DESC, id DESC
        """,
        (usuario_id,),
    )
    mensajes_no_leidos = cursor.fetchall()
    conn.close()
    mensajes_filtrados = [
        fila for fila in mensajes_no_leidos if ("mensaje_no_leido", str(fila["id"])) not in descartes_hoy
    ]
    mensajes_total = len(mensajes_filtrados)
    mensajes_recientes = [
        {
            "id": fila["id"],
            "asunto": fila["asunto"],
            "remitente": fila["remitente"],
        }
        for fila in mensajes_filtrados[:5]
    ]

    return {
        "vencidas": vencidas,
        "por_vencer": por_vencer,
        "metas_rezagadas": metas_rezagadas,
        "eventos_hoy": eventos_hoy,
        "eventos_proximos": eventos_proximos,
        "mensajes_no_leidos_total": mensajes_total,
        "mensajes_no_leidos_recientes": mensajes_recientes,
    }


def obtener_kpis_usuario(usuario_id):
    conn = get_connection()
    cursor = conn.cursor()
    hoy = datetime.now().date().strftime("%Y-%m-%d")

    cursor.execute("SELECT COUNT(*) FROM metas WHERE user_id = ?", (usuario_id,))
    metas_total = int(cursor.fetchone()[0])

    cursor.execute(
        "SELECT COUNT(*) FROM calendario_eventos WHERE user_id = ? AND fecha_evento = ?",
        (usuario_id, hoy),
    )
    eventos_hoy = int(cursor.fetchone()[0])

    cursor.execute(
        "SELECT COUNT(*) FROM mensajes WHERE user_id = ? AND carpeta = 'entrada' AND leido = 0",
        (usuario_id,),
    )
    mensajes_no_leidos = int(cursor.fetchone()[0])

    conn.close()
    notis = obtener_notificaciones_usuario(usuario_id)
    return {
        "metas_total": metas_total,
        "eventos_hoy": eventos_hoy,
        "mensajes_no_leidos": mensajes_no_leidos,
        "alertas_total": notificaciones_compose.total_alertas(notis),
    }


def obtener_actividad_reciente_usuario(usuario_id, limite=10):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT modulo, accion, entidad, entidad_id, creado_en, detalle
        FROM audit_log
        WHERE actor_user_id = ?
        ORDER BY datetime(creado_en) DESC, id DESC
        LIMIT ?
        """,
        (usuario_id, max(1, int(limite))),
    )
    filas = cursor.fetchall()
    conn.close()
    return filas


@app.context_processor
def inyectar_notificaciones():
    if not usuario_autenticado():
        return {"noti_count": 0}

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        return {"noti_count": 0}

    notis = obtener_notificaciones_usuario(usuario_id)
    return {"noti_count": notificaciones_compose.total_alertas(notis)}


@app.context_processor
def inyectar_auth_social():
    return {
        "facebook_login_enabled": facebook_configurado(),
        "google_login_enabled": google_configurado(),
        "microsoft_login_enabled": microsoft_configurado(),
        "whatsapp_registro_enabled": whatsapp_configurado(),
    }


@app.context_processor
def inyectar_seguridad_plantillas():
    return {
        "csrf_token": token_csrf_actual(),
        "current_role": obtener_rol_usuario_actual(),
    }


@app.context_processor
def inyectar_kpis_top():
    if not usuario_autenticado():
        return {"kpis_top": None}
    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        return {"kpis_top": None}
    return {"kpis_top": obtener_kpis_usuario(usuario_id)}


def notificaciones_en_mantenimiento():
    return os.path.exists(MAINTENANCE_NOTIFICATIONS_FLAG)


def set_mantenimiento_notificaciones(activo, actor_user_id=None):
    os.makedirs(os.path.join(BASE_DIR, "logs"), exist_ok=True)
    if activo:
        with open(MAINTENANCE_NOTIFICATIONS_FLAG, "w", encoding="utf-8") as fh:
            fh.write(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    else:
        try:
            os.remove(MAINTENANCE_NOTIFICATIONS_FLAG)
        except FileNotFoundError:
            pass
    if actor_user_id:
        registrar_auditoria(
            actor_user_id,
            "admin",
            "mantenimiento_notificaciones",
            "sistema",
            None,
            {"activo": bool(activo)},
        )


def login():
    if usuario_autenticado():
        return redirect(url_for("metas_bp.inicio_metas"))

    remembered_token = (request.cookies.get("recordar_token") or "").strip()
    remembered_session = validar_token_recordar(remembered_token) if remembered_token else None
    if remembered_session:
        iniciar_sesion_usuario(
            remembered_session["id"],
            remembered_session["username"],
            remembered_session.get("rol"),
        )
        session.permanent = True
        return redirect(url_for("metas_bp.inicio_metas"))

    requested_user = (request.args.get("usuario") or "").strip()
    remembered_user = requested_user or (request.cookies.get("recordar_usuario") or "").strip()

    if request.method == "POST":
        ip_cliente = obtener_ip_cliente()
        usuario_ingresado = (
            request.form.get("username")
            or request.form.get("usuario")
            or ""
        ).strip()
        accion_login_usuario = _accion_rate_login_usuario(usuario_ingresado)
        bloqueado_ip, segundos_ip = verificar_rate_limit("login", ip_cliente)
        bloqueado_user, segundos_user = (False, 0)
        if accion_login_usuario:
            bloqueado_user, segundos_user = verificar_rate_limit(accion_login_usuario, ip_cliente)
        bloqueado = bloqueado_ip or bloqueado_user
        segundos_restantes = max(segundos_ip, segundos_user)
        recordar_sesion = request.form.get("recordar_sesion") == "1"
        if bloqueado:
            flash(
                "Demasiados intentos de inicio de sesion. "
                f"Intenta de nuevo en {_formato_tiempo_restante(segundos_restantes)}.",
                "warning",
            )
            return render_template(
                "auth/login.html",
                login_usuario=usuario_ingresado or remembered_user,
                recordar_sesion=recordar_sesion or bool(remembered_user),
            )

        usuario = usuario_ingresado
        contrasena = request.form.get("password") or request.form.get("contrasena") or ""

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, username, password, rol FROM usuarios WHERE username = ?",
            (usuario,),
        )
        user = cursor.fetchone()
        acceso_ok = False

        if user:
            acceso_ok, requiere_upgrade = verificar_contrasena(user["password"], contrasena)
            if acceso_ok and requiere_upgrade:
                cursor.execute(
                    "UPDATE usuarios SET password = ? WHERE id = ?",
                    (hash_contrasena(contrasena), user["id"]),
                )
                conn.commit()

        conn.close()

        if user and acceso_ok:
            limpiar_rate_limit("login", ip_cliente)
            if accion_login_usuario:
                limpiar_rate_limit(accion_login_usuario, ip_cliente)
            registrar_login_attempt(usuario_ingresado, ip_cliente, True)
            if _admin_2fa_requerido_para_user(user):
                ok_2fa, err_2fa = iniciar_admin_2fa_pendiente(user)
                if ok_2fa:
                    flash("Te enviamos un codigo 2FA al correo admin.", "info")
                    return redirect(url_for("auth_bp.admin_2fa"))
                flash(f"No se pudo activar 2FA admin: {err_2fa}", "warning")
            iniciar_sesion_usuario(user["id"], user["username"], user["rol"])
            session.permanent = recordar_sesion
            respuesta = make_response(redirect(url_for("metas_bp.inicio_metas")))
            if recordar_sesion:
                token = generar_token_recordar(user["id"], user["username"])
                respuesta.set_cookie(
                    "recordar_usuario",
                    user["username"],
                    max_age=60 * 60 * 24 * REMEMBER_ME_DAYS,
                    samesite=COOKIE_SAMESITE,
                    secure=COOKIE_SECURE,
                    httponly=True,
                )
                respuesta.set_cookie(
                    "recordar_token",
                    token,
                    max_age=60 * 60 * 24 * REMEMBER_ME_DAYS,
                    samesite=COOKIE_SAMESITE,
                    secure=COOKIE_SECURE,
                    httponly=True,
                )
            else:
                respuesta.delete_cookie("recordar_usuario")
                respuesta.delete_cookie("recordar_token")
            return respuesta

        registrar_login_attempt(usuario_ingresado, ip_cliente, False)
        info_ip = registrar_intento_rate_limit("login", ip_cliente) or {}
        info_user = registrar_intento_rate_limit(accion_login_usuario, ip_cliente) if accion_login_usuario else {}
        if user and (info_ip.get("bloqueado_hasta") or (info_user or {}).get("bloqueado_hasta")):
            registrar_auditoria(
                user["id"],
                "auth",
                "login_bloqueado",
                "usuario",
                user["id"],
                {"ip": ip_cliente, "usuario": usuario_ingresado},
            )
        flash("Usuario o contrasena incorrecta.", "danger")
        return render_template(
            "auth/login.html",
            login_usuario=usuario_ingresado or remembered_user,
            recordar_sesion=recordar_sesion or bool(remembered_user),
        )

    return render_template(
        "auth/login.html",
        login_usuario=remembered_user,
        recordar_sesion=bool(remembered_user),
    )


def admin_2fa():
    pending = session.get("admin_2fa_pending")
    if not pending:
        flash("No hay verificacion 2FA pendiente.", "warning")
        return redirect(url_for("auth_bp.login"))

    if request.method == "POST":
        codigo = (request.form.get("codigo") or "").strip()
        ok, data = validar_admin_2fa_pending(codigo)
        if not ok:
            flash(str(data), "danger")
            return render_template("auth/admin_2fa.html")

        session.pop("admin_2fa_pending", None)
        iniciar_sesion_usuario(data["id"], data["username"], data["rol"])
        session.permanent = True
        flash("2FA validado. Bienvenido.", "success")
        return redirect(url_for("admin_dashboard"))

    return render_template("auth/admin_2fa.html")


def facebook_start():
    if not facebook_configurado():
        flash("Facebook login no esta configurado.", "warning")
        return redirect(url_for("auth_bp.login"))

    state = secrets.token_urlsafe(24)
    session["fb_oauth_state"] = state
    redirect_uri = url_for("auth_bp.facebook_callback", _external=True)
    params = urlparse.urlencode(
        {
            "client_id": FACEBOOK_APP_ID,
            "redirect_uri": redirect_uri,
            "state": state,
            "scope": FACEBOOK_SCOPE,
            "response_type": "code",
        }
    )
    return redirect(f"https://www.facebook.com/{FACEBOOK_GRAPH_VERSION}/dialog/oauth?{params}")


def facebook_callback():
    if not facebook_configurado():
        flash("Facebook login no esta configurado.", "warning")
        return redirect(url_for("auth_bp.login"))

    state_esperado = session.pop("fb_oauth_state", None)
    state = request.args.get("state", "")
    code = request.args.get("code", "")
    if not state_esperado or state != state_esperado or not code:
        flash("No se pudo validar el login de Facebook.", "danger")
        return redirect(url_for("auth_bp.login"))

    redirect_uri = url_for("auth_bp.facebook_callback", _external=True)
    token_url = (
        f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/oauth/access_token?"
        + urlparse.urlencode(
            {
                "client_id": FACEBOOK_APP_ID,
                "client_secret": FACEBOOK_APP_SECRET,
                "redirect_uri": redirect_uri,
                "code": code,
            }
        )
    )
    try:
        with urlrequest.urlopen(token_url, timeout=15) as resp:
            token_data = json.loads(resp.read().decode("utf-8"))
        access_token = token_data.get("access_token", "")
        if not access_token:
            raise RuntimeError("No se recibio access token.")

        me_url = (
            f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/me?"
            + urlparse.urlencode(
                {
                    "fields": "id,name,email",
                    "access_token": access_token,
                }
            )
        )
        with urlrequest.urlopen(me_url, timeout=15) as resp:
            perfil = json.loads(resp.read().decode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        app.logger.exception("Error en OAuth de Facebook")
        flash(f"No se pudo completar Facebook login: {exc}", "danger")
        return redirect(url_for("auth_bp.login"))

    fb_id = str(perfil.get("id", "")).strip()
    fb_name = str(perfil.get("name", "")).strip()
    fb_email = str(perfil.get("email", "")).strip().lower()
    if not fb_id:
        flash("Facebook no devolvio un identificador valido.", "danger")
        return redirect(url_for("auth_bp.login"))

    user_id, username = crear_o_obtener_usuario_facebook(fb_id, fb_name, fb_email)
    iniciar_sesion_usuario(user_id, username)
    flash("Sesion iniciada con Facebook.", "success")
    return redirect(url_for("metas_bp.inicio_metas"))


def google_start():
    if not google_configurado():
        flash("Google login no esta configurado (faltan CLIENT_ID o CLIENT_SECRET).", "warning")
        return redirect(url_for("auth_bp.login"))

    state = secrets.token_urlsafe(24)
    session["google_oauth_state"] = state
    redirect_uri = _oauth_redirect_uri("auth_bp.google_callback", GOOGLE_REDIRECT_URI)
    params = urlparse.urlencode(
        {
            "client_id": GOOGLE_CLIENT_ID,
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": GOOGLE_SCOPE,
            "state": state,
            "prompt": "select_account",
        }
    )
    return redirect(f"https://accounts.google.com/o/oauth2/v2/auth?{params}")


def google_callback():
    if not google_configurado():
        flash("Google login no esta configurado (faltan CLIENT_ID o CLIENT_SECRET).", "warning")
        return redirect(url_for("auth_bp.login"))

    state_esperado = session.pop("google_oauth_state", None)
    state = request.args.get("state", "")
    code = request.args.get("code", "")
    if not state_esperado or state != state_esperado or not code:
        flash("No se pudo validar el login de Google.", "danger")
        return redirect(url_for("auth_bp.login"))

    redirect_uri = _oauth_redirect_uri("auth_bp.google_callback", GOOGLE_REDIRECT_URI)
    try:
        token_payload = urlparse.urlencode(
            {
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "code": code,
                "grant_type": "authorization_code",
                "redirect_uri": redirect_uri,
            }
        ).encode("utf-8")
        token_req = urlrequest.Request(
            "https://oauth2.googleapis.com/token",
            data=token_payload,
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            method="POST",
        )
        with urlrequest.urlopen(token_req, timeout=15) as resp:
            token_data = json.loads(resp.read().decode("utf-8"))

        access_token = token_data.get("access_token", "")
        if not access_token:
            raise RuntimeError("No se recibio access token de Google.")

        me_req = urlrequest.Request(
            "https://openidconnect.googleapis.com/v1/userinfo",
            headers={"Authorization": f"Bearer {access_token}"},
            method="GET",
        )
        with urlrequest.urlopen(me_req, timeout=15) as resp:
            perfil = json.loads(resp.read().decode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        app.logger.exception("Error en OAuth de Google")
        flash(f"No se pudo completar Google login: {exc}", "danger")
        return redirect(url_for("auth_bp.login"))

    oauth_id = str(perfil.get("sub", "")).strip()
    nombre = str(perfil.get("name", "")).strip()
    email = str(perfil.get("email", "")).strip().lower()
    if not oauth_id:
        flash("Google no devolvio un identificador valido.", "danger")
        return redirect(url_for("auth_bp.login"))

    user_id, username = crear_o_obtener_usuario_oauth("google", oauth_id, nombre, email)
    iniciar_sesion_usuario(user_id, username)
    flash("Sesion iniciada con Google.", "success")
    return redirect(url_for("metas_bp.inicio_metas"))


def microsoft_start():
    if not microsoft_configurado():
        flash("Microsoft login no esta configurado.", "warning")
        return redirect(url_for("auth_bp.login"))

    state = secrets.token_urlsafe(24)
    session["microsoft_oauth_state"] = state
    redirect_uri = url_for("auth_bp.microsoft_callback", _external=True)
    params = urlparse.urlencode(
        {
            "client_id": MICROSOFT_CLIENT_ID,
            "response_type": "code",
            "redirect_uri": redirect_uri,
            "scope": MICROSOFT_SCOPE,
            "state": state,
            "prompt": "select_account",
        }
    )
    return redirect(
        f"https://login.microsoftonline.com/{MICROSOFT_TENANT}/oauth2/v2.0/authorize?{params}"
    )


def microsoft_callback():
    if not microsoft_configurado():
        flash("Microsoft login no esta configurado.", "warning")
        return redirect(url_for("auth_bp.login"))

    state_esperado = session.pop("microsoft_oauth_state", None)
    state = request.args.get("state", "")
    code = request.args.get("code", "")
    if not state_esperado or state != state_esperado or not code:
        flash("No se pudo validar el login de Microsoft.", "danger")
        return redirect(url_for("auth_bp.login"))

    redirect_uri = url_for("auth_bp.microsoft_callback", _external=True)
    try:
        token_payload = urlparse.urlencode(
            {
                "client_id": MICROSOFT_CLIENT_ID,
                "client_secret": MICROSOFT_CLIENT_SECRET,
                "code": code,
                "grant_type": "authorization_code",
                "redirect_uri": redirect_uri,
                "scope": MICROSOFT_SCOPE,
            }
        ).encode("utf-8")
        token_req = urlrequest.Request(
            f"https://login.microsoftonline.com/{MICROSOFT_TENANT}/oauth2/v2.0/token",
            data=token_payload,
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            method="POST",
        )
        with urlrequest.urlopen(token_req, timeout=15) as resp:
            token_data = json.loads(resp.read().decode("utf-8"))

        access_token = token_data.get("access_token", "")
        if not access_token:
            raise RuntimeError("No se recibio access token de Microsoft.")

        me_req = urlrequest.Request(
            "https://graph.microsoft.com/v1.0/me",
            headers={"Authorization": f"Bearer {access_token}"},
            method="GET",
        )
        with urlrequest.urlopen(me_req, timeout=15) as resp:
            perfil = json.loads(resp.read().decode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        app.logger.exception("Error en OAuth de Microsoft")
        flash(f"No se pudo completar Microsoft login: {exc}", "danger")
        return redirect(url_for("auth_bp.login"))

    oauth_id = str(perfil.get("id", "")).strip()
    nombre = str(perfil.get("displayName", "")).strip()
    email = str(perfil.get("mail") or perfil.get("userPrincipalName") or "").strip().lower()
    if not oauth_id:
        flash("Microsoft no devolvio un identificador valido.", "danger")
        return redirect(url_for("auth_bp.login"))

    user_id, username = crear_o_obtener_usuario_oauth("microsoft", oauth_id, nombre, email)
    iniciar_sesion_usuario(user_id, username)
    flash("Sesion iniciada con Microsoft.", "success")
    return redirect(url_for("metas_bp.inicio_metas"))


def registro():
    if request.method == "POST":
        usuario = request.form["usuario"].strip()
        email = request.form["email"].strip().lower()
        telefono = limpiar_telefono(request.form.get("telefono", ""))
        contrasena = request.form["contrasena"]

        if not email:
            flash("El email es obligatorio.", "warning")
            return redirect(url_for("auth_bp.registro"))

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT 1 FROM usuarios WHERE username = ?", (usuario,))
        if cursor.fetchone():
            conn.close()
            flash("El usuario ya existe.", "warning")
            return redirect(url_for("auth_bp.registro"))

        cursor.execute("SELECT 1 FROM usuarios WHERE email = ?", (email,))
        if cursor.fetchone():
            conn.close()
            flash("El email ya esta en uso.", "warning")
            return redirect(url_for("auth_bp.registro"))

        if telefono and len(telefono) < 10:
            conn.close()
            flash("Ingresa un telefono valido en formato internacional.", "warning")
            return redirect(url_for("auth_bp.registro"))

        if telefono:
            cursor.execute("SELECT 1 FROM usuarios WHERE telefono = ?", (telefono,))
            if cursor.fetchone():
                conn.close()
                flash("Ese telefono ya esta registrado.", "warning")
                return redirect(url_for("auth_bp.registro"))

        cursor.execute(
            "INSERT INTO usuarios (username, password, email, telefono, rol) VALUES (?, ?, ?, ?, ?)",
            (usuario, hash_contrasena(contrasena), email, telefono or None, "editor"),
        )
        conn.commit()
        conn.close()

        flash("Usuario registrado correctamente.", "success")
        return redirect(url_for("auth_bp.login"))

    return render_template("auth/registro.html")


def registro_whatsapp():
    whatsapp_registro_habilitado = False
    if not whatsapp_registro_habilitado:
        flash("Registro por WhatsApp deshabilitado temporalmente.", "info")
        return redirect(url_for("auth_bp.registro"))

    ip_cliente = obtener_ip_cliente()
    accion = request.form.get("accion", "").strip()
    usuario = request.form.get("usuario", "").strip()
    telefono = limpiar_telefono(request.form.get("telefono", ""))
    contrasena = request.form.get("contrasena", "")
    codigo = request.form.get("codigo", "").strip()

    if accion == "enviar":
        bloqueado, segundos_restantes = verificar_rate_limit("registro_wa", ip_cliente)
        if bloqueado:
            flash(
                "Demasiados intentos para enviar codigo de WhatsApp. "
                f"Intenta de nuevo en {_formato_tiempo_restante(segundos_restantes)}.",
                "warning",
            )
            return redirect(url_for("auth_bp.registro_whatsapp"))

        if not whatsapp_configurado():
            flash("WhatsApp no esta configurado en el servidor.", "danger")
            return redirect(url_for("auth_bp.registro_whatsapp"))
        if not usuario or len(usuario) < 3:
            flash("Usuario invalido. Minimo 3 caracteres.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))
        if not telefono or len(telefono) < 10:
            flash("Telefono invalido. Usa formato internacional con codigo de pais.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))
        if len(contrasena) < 4:
            flash("La contrasena debe tener al menos 4 caracteres.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT 1 FROM usuarios WHERE username = ?", (usuario,))
        if cursor.fetchone():
            conn.close()
            flash("El usuario ya existe.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))
        cursor.execute("SELECT 1 FROM usuarios WHERE telefono = ?", (telefono,))
        if cursor.fetchone():
            conn.close()
            flash("Ese telefono ya esta registrado.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))
        conn.close()

        codigo_generado = f"{secrets.randbelow(1000000):06d}"
        expira = (datetime.now() + timedelta(minutes=10)).strftime("%Y-%m-%d %H:%M:%S")
        session["wa_registro"] = {
            "usuario": usuario,
            "telefono": telefono,
            "contrasena_hash": hash_contrasena(contrasena),
            "codigo": codigo_generado,
            "expira": expira,
        }

        try:
            enviar_whatsapp_codigo(telefono, codigo_generado)
        except Exception as exc:  # noqa: BLE001
            app.logger.exception("Error enviando codigo WhatsApp de registro")
            flash(f"No se pudo enviar el codigo por WhatsApp: {exc}", "danger")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        registrar_intento_rate_limit("registro_wa", ip_cliente)
        flash("Codigo enviado por WhatsApp. Revisa tu telefono.", "success")
        return redirect(url_for("auth_bp.registro_whatsapp"))

    if accion == "verificar":
        data = session.get("wa_registro")
        if not data:
            flash("No hay una solicitud de registro activa. Envia un codigo primero.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        if codigo != data.get("codigo", ""):
            flash("Codigo incorrecto.", "danger")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        fecha_expira = parsear_fecha_hora(data.get("expira", ""))
        if fecha_expira is None or fecha_expira < datetime.now():
            session.pop("wa_registro", None)
            flash("El codigo expiro. Solicita uno nuevo.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT 1 FROM usuarios WHERE username = ?", (data["usuario"],))
        if cursor.fetchone():
            conn.close()
            flash("El usuario ya existe.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        cursor.execute("SELECT 1 FROM usuarios WHERE telefono = ?", (data["telefono"],))
        if cursor.fetchone():
            conn.close()
            flash("Ese telefono ya esta registrado.", "warning")
            return redirect(url_for("auth_bp.registro_whatsapp"))

        cursor.execute(
            "INSERT INTO usuarios (username, password, telefono, rol) VALUES (?, ?, ?, ?)",
            (data["usuario"], data["contrasena_hash"], data["telefono"], "editor"),
        )
        conn.commit()
        user_id = cursor.lastrowid
        conn.close()

        session.pop("wa_registro", None)
        limpiar_rate_limit("registro_wa", ip_cliente)
        iniciar_sesion_usuario(user_id, data["usuario"])
        flash("Cuenta creada y verificada por WhatsApp.", "success")
        return redirect(url_for("metas_bp.inicio_metas"))

    flash("Accion de WhatsApp no valida.", "warning")
    return redirect(url_for("auth_bp.registro_whatsapp"))


def recuperar_cuenta():
    correo_enviado = False
    error_envio = False
    destino_mask = None
    error_detalle = None
    proveedor_usado = None

    if request.method == "POST":
        ip_cliente = obtener_ip_cliente()
        bloqueado, segundos_restantes = verificar_rate_limit("recuperar", ip_cliente)
        if bloqueado:
            flash(
                "Demasiadas solicitudes de recuperacion. "
                f"Intenta de nuevo en {_formato_tiempo_restante(segundos_restantes)}.",
                "warning",
            )
            return render_template(
                "auth/recuperar.html",
                correo_config_ok=proveedor_correo_configurado(),
                correo_provider=nombre_ruta_proveedores(),
            )

        registrar_intento_rate_limit("recuperar", ip_cliente)
        identificador = request.form["identificador"].strip().lower()

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, email FROM usuarios WHERE lower(username) = ? OR lower(email) = ?",
            (identificador, identificador),
        )
        user = cursor.fetchone()

        if user:
            email_destino = (user["email"] or "").strip()
            if not proveedor_correo_configurado() or not email_destino:
                error_envio = True
                if not email_destino:
                    error_detalle = "El usuario no tiene email configurado."
            else:
                codigo = f"{secrets.randbelow(1000000):06d}"
                expira = (datetime.now() + timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M:%S")
                try:
                    proveedor_usado = enviar_email_recuperacion(email_destino, codigo)
                    correo_enviado = bool(proveedor_usado)
                except Exception as exc:
                    app.logger.exception("Error enviando codigo de recuperacion")
                    correo_enviado = False
                    error_detalle = str(exc)

                if correo_enviado:
                    cursor.execute(
                        "UPDATE usuarios SET reset_token = ?, reset_expira = ? WHERE id = ?",
                        (codigo, expira, user["id"]),
                    )
                    conn.commit()
                    destino_mask = enmascarar_email(email_destino)
                else:
                    error_envio = True

        conn.close()

        if user and correo_enviado:
            flash(
                f"Si la cuenta existe, enviamos un codigo a {destino_mask} via {proveedor_usado}.",
                "success",
            )
        elif user and error_envio:
            flash(
                "No se pudo enviar el codigo por correo. Revisa el proveedor y el email del usuario.",
                "danger",
            )
            if app.debug and error_detalle:
                flash(f"Detalle tecnico: {error_detalle}", "warning")
        else:
            flash(
                "Si la cuenta existe, se genero un codigo de recuperacion valido por 15 minutos.",
                "info",
            )

    return render_template(
        "auth/recuperar.html",
        correo_config_ok=proveedor_correo_configurado(),
        correo_provider=nombre_ruta_proveedores(),
    )


def restablecer_contrasena():
    if request.method == "GET":
        return render_template("auth/restablecer.html")

    identificador = request.form["identificador"].strip().lower()
    codigo = request.form["codigo"].strip()
    contrasena = request.form["contrasena"]
    confirmar = request.form["confirmar_contrasena"]

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, reset_token, reset_expira FROM usuarios WHERE lower(username) = ? OR lower(email) = ?",
        (identificador, identificador),
    )
    user = cursor.fetchone()

    if user is None:
        conn.close()
        flash("Usuario o email no valido.", "danger")
        return redirect(url_for("auth_bp.recuperar_cuenta"))

    if (user["reset_token"] or "").strip() != codigo:
        conn.close()
        flash("Codigo de recuperacion incorrecto.", "danger")
        return redirect(url_for("auth_bp.restablecer_contrasena"))

    fecha_expira = parsear_fecha_hora(user["reset_expira"])
    if fecha_expira is None or fecha_expira < datetime.now():
        cursor.execute(
            "UPDATE usuarios SET reset_token = NULL, reset_expira = NULL WHERE id = ?",
            (user["id"],),
        )
        conn.commit()
        conn.close()
        flash("El codigo de recuperacion expiro. Genera uno nuevo.", "warning")
        return redirect(url_for("auth_bp.recuperar_cuenta"))

    if len(contrasena) < 4:
        conn.close()
        flash("La contrasena debe tener al menos 4 caracteres.", "warning")
        return redirect(url_for("auth_bp.restablecer_contrasena"))

    if contrasena != confirmar:
        conn.close()
        flash("Las contrasenas no coinciden.", "warning")
        return redirect(url_for("auth_bp.restablecer_contrasena"))

    cursor.execute(
        "UPDATE usuarios SET password = ?, reset_token = NULL, reset_expira = NULL WHERE id = ?",
        (hash_contrasena(contrasena), user["id"]),
    )
    conn.commit()
    conn.close()

    flash("Contrasena actualizada correctamente.", "success")
    return redirect(url_for("auth_bp.login"))


def logout():
    switch_user = (request.args.get("switch_user") or "").strip()
    clear_remember = request.args.get("clear_remember", "0").strip() in ("1", "true", "True")
    session.pop("usuario", None)
    session.pop("usuario_id", None)
    session.pop("usuario_rol", None)
    session.pop("csrf_token", None)
    flash("Sesion cerrada.", "info")
    respuesta = make_response(redirect(url_for("auth_bp.login")))
    if switch_user:
        respuesta = make_response(redirect(url_for("auth_bp.login", usuario=switch_user)))
        respuesta.set_cookie(
            "recordar_usuario",
            switch_user,
            max_age=60 * 60 * 24 * REMEMBER_ME_DAYS,
            samesite=COOKIE_SAMESITE,
            secure=COOKIE_SECURE,
            httponly=True,
        )
    if clear_remember:
        respuesta.delete_cookie("recordar_usuario")
    respuesta.delete_cookie("recordar_token")
    return respuesta


def _anio_mes_solicitado():
    hoy = datetime.now()
    try:
        anio = int(request.args.get("anio", hoy.year))
        mes = int(request.args.get("mes", hoy.month))
    except ValueError:
        return hoy.year, hoy.month

    if anio < 2000 or anio > 2100:
        anio = hoy.year
    if mes < 1 or mes > 12:
        mes = hoy.month
    return anio, mes


def _filtros_calendario():
    return {
        "grupo": request.args.get("grupo", "").strip(),
        "lugar": request.args.get("lugar", "").strip(),
        "tipo": request.args.get("tipo", "").strip(),
    }


def _contexto_redireccion_calendario(anio_default=None, mes_default=None):
    return calendario_logic.construir_contexto_redireccion_calendario(request, anio_default, mes_default)


def _obtener_opciones_filtros_calendario(usuario_id):
    return calendario_queries.obtener_opciones_filtros_calendario(get_connection, usuario_id)


def _obtener_eventos_calendario(usuario_id, anio, mes, filtros):
    return calendario_queries.obtener_eventos_calendario(get_connection, usuario_id, anio, mes, filtros)


def _matriz_calendario(anio, mes, eventos):
    return calendario_queries.matriz_calendario(anio, mes, eventos)


def calendario():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    anio, mes = _anio_mes_solicitado()
    filtros = _filtros_calendario()
    eventos = _obtener_eventos_calendario(usuario_id, anio, mes, filtros)
    semanas = _matriz_calendario(anio, mes, eventos)
    opciones = _obtener_opciones_filtros_calendario(usuario_id)
    try:
        page_eventos = max(1, int(request.args.get("page_eventos", "1")))
    except ValueError:
        page_eventos = 1
    page_size_eventos = 8
    total_eventos = len(eventos)
    total_paginas_eventos = max(1, (total_eventos + page_size_eventos - 1) // page_size_eventos)
    if page_eventos > total_paginas_eventos:
        page_eventos = total_paginas_eventos
    ini = (page_eventos - 1) * page_size_eventos
    fin = ini + page_size_eventos
    eventos_pagina = eventos[ini:fin]

    fecha_ref = datetime(anio, mes, 1)
    fecha_prev = (fecha_ref - timedelta(days=1)).replace(day=1)
    fecha_next = (fecha_ref + timedelta(days=32)).replace(day=1)

    return render_template(
        "calendario/calendario.html",
        anio=anio,
        mes=mes,
        mes_nombre=MESES_ES[mes - 1],
        dias_corto=DIAS_ES_CORTO,
        semanas=semanas,
        filtros=filtros,
        opciones=opciones,
        eventos=eventos,
        eventos_pagina=eventos_pagina,
        page_eventos=page_eventos,
        total_paginas_eventos=total_paginas_eventos,
        total_eventos=total_eventos,
        prev_anio=fecha_prev.year,
        prev_mes=fecha_prev.month,
        next_anio=fecha_next.year,
        next_mes=fecha_next.month,
        hoy_anio=datetime.now().year,
        hoy_mes=datetime.now().month,
    )


def crear_evento():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    titulo = request.form.get("titulo", "").strip()
    fecha_evento = request.form.get("fecha_evento", "").strip()
    hora_evento = request.form.get("hora_evento", "").strip()
    grupo = request.form.get("grupo", "").strip()
    lugar = request.form.get("lugar", "").strip()
    tipo = request.form.get("tipo", "").strip()
    descripcion = request.form.get("descripcion", "").strip()

    fecha_obj = parsear_fecha_limite(fecha_evento)
    if not titulo or fecha_obj is None:
        flash("Debes ingresar al menos titulo y fecha valida.", "warning")
        return redirect(url_for("calendario_bp.calendario"))

    if hora_evento and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", hora_evento):
        flash("La hora debe tener formato HH:MM.", "warning")
        return redirect(url_for("calendario_bp.calendario", anio=fecha_obj.year, mes=fecha_obj.month))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO calendario_eventos
        (user_id, titulo, fecha_evento, hora_evento, grupo, lugar, tipo, descripcion, creado_en)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            usuario_id,
            titulo,
            fecha_obj.strftime("%Y-%m-%d"),
            hora_evento or None,
            grupo or None,
            lugar or None,
            tipo or None,
            descripcion or None,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    )
    evento_id = cursor.lastrowid
    conn.commit()
    conn.close()
    registrar_auditoria(
        usuario_id,
        "calendario",
        "crear",
        "evento",
        evento_id,
        {"titulo": titulo, "fecha_evento": fecha_obj.strftime("%Y-%m-%d")},
    )

    flash("Evento creado.", "success")
    return redirect(url_for("calendario_bp.calendario", anio=fecha_obj.year, mes=fecha_obj.month))


def editar_evento(evento_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    contexto = _contexto_redireccion_calendario()

    titulo = request.form.get("titulo", "").strip()
    fecha_evento = request.form.get("fecha_evento", "").strip()
    hora_evento = request.form.get("hora_evento", "").strip()
    grupo = request.form.get("grupo", "").strip()
    lugar = request.form.get("lugar", "").strip()
    tipo = request.form.get("tipo", "").strip()
    descripcion = request.form.get("descripcion", "").strip()

    fecha_obj = parsear_fecha_limite(fecha_evento)
    if not titulo or fecha_obj is None:
        flash("Debes ingresar titulo y fecha valida para editar el evento.", "warning")
        return redirect(url_for("calendario_bp.calendario", **contexto))

    if hora_evento and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", hora_evento):
        flash("La hora debe tener formato HH:MM.", "warning")
        return redirect(url_for("calendario_bp.calendario", **contexto))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        UPDATE calendario_eventos
        SET titulo = ?, fecha_evento = ?, hora_evento = ?, grupo = ?, lugar = ?, tipo = ?, descripcion = ?
        WHERE id = ? AND user_id = ?
        """,
        (
            titulo,
            fecha_obj.strftime("%Y-%m-%d"),
            hora_evento or None,
            grupo or None,
            lugar or None,
            tipo or None,
            descripcion or None,
            evento_id,
            usuario_id,
        ),
    )
    conn.commit()
    actualizados = cursor.rowcount
    conn.close()

    if actualizados == 0:
        flash("No se pudo editar el evento (no existe o no te pertenece).", "warning")
        return redirect(url_for("calendario_bp.calendario", **contexto))

    registrar_auditoria(
        usuario_id,
        "calendario",
        "editar",
        "evento",
        evento_id,
        {"titulo": titulo, "fecha_evento": fecha_obj.strftime("%Y-%m-%d")},
    )
    flash("Evento actualizado.", "success")
    return redirect(
        url_for(
            "calendario_bp.calendario",
            anio=fecha_obj.year,
            mes=fecha_obj.month,
            grupo=contexto["grupo"],
            lugar=contexto["lugar"],
            tipo=contexto["tipo"],
            page_eventos=contexto["page_eventos"],
        )
    )


def eliminar_evento(evento_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    contexto = _contexto_redireccion_calendario()

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "DELETE FROM calendario_eventos WHERE id = ? AND user_id = ?",
        (evento_id, usuario_id),
    )
    conn.commit()
    eliminados = cursor.rowcount
    conn.close()

    if eliminados == 0:
        flash("No se pudo eliminar el evento (no existe o no te pertenece).", "warning")
    else:
        registrar_auditoria(usuario_id, "calendario", "eliminar", "evento", evento_id, {})
        flash("Evento eliminado.", "info")

    return redirect(url_for("calendario_bp.calendario", **contexto))


def exportar_eventos_csv():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    anio, mes = _anio_mes_solicitado()
    filtros = _filtros_calendario()
    eventos = _obtener_eventos_calendario(usuario_id, anio, mes, filtros)

    salida = StringIO()
    writer = csv.writer(salida)
    writer.writerow(["fecha", "hora", "titulo", "grupo", "lugar", "tipo", "descripcion"])
    for fila in eventos:
        writer.writerow(
            [
                fila["fecha_evento"],
                fila["hora_evento"] or "",
                fila["titulo"],
                fila["grupo"] or "",
                fila["lugar"] or "",
                fila["tipo"] or "",
                fila["descripcion"] or "",
            ]
        )

    contenido = salida.getvalue()
    nombre = f"calendario_{anio:04d}_{mes:02d}.csv"
    return Response(
        contenido,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


def exportar_eventos_pdf():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    anio, mes = _anio_mes_solicitado()
    filtros = _filtros_calendario()
    if isinstance(filtros, dict):
        filtro_grupo = filtros.get("grupo", "")
        filtro_lugar = filtros.get("lugar", "")
        filtro_tipo = filtros.get("tipo", "")
    else:
        filtro_grupo = getattr(filtros, "grupo", "")
        filtro_lugar = getattr(filtros, "lugar", "")
        filtro_tipo = getattr(filtros, "tipo", "")
    eventos = _obtener_eventos_calendario(usuario_id, anio, mes, filtros)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except Exception:
        flash("Para exportar PDF instala reportlab (pip install reportlab).", "warning")
        return redirect(
            url_for(
                "calendario_bp.calendario",
                anio=anio,
                mes=mes,
                grupo=filtro_grupo,
                lugar=filtro_lugar,
                tipo=filtro_tipo,
            )
        )

    buff = BytesIO()
    c = canvas.Canvas(buff, pagesize=A4)
    w, h = A4
    y = h - 40

    c.setFont("Helvetica-Bold", 13)
    c.drawString(36, y, f"Calendario {anio:04d}-{mes:02d}")
    y -= 16
    c.setFont("Helvetica", 9)
    c.drawString(
        36,
        y,
        f"Filtros: grupo={filtro_grupo or '-'} | lugar={filtro_lugar or '-'} | tipo={filtro_tipo or '-'}",
    )
    y -= 20

    c.setFont("Helvetica-Bold", 9)
    c.drawString(36, y, "Fecha/Hora")
    c.drawString(130, y, "Titulo")
    c.drawString(305, y, "Grupo")
    c.drawString(390, y, "Lugar")
    c.drawString(470, y, "Tipo")
    y -= 12
    c.setFont("Helvetica", 8)

    for fila in eventos:
        if y < 48:
            c.showPage()
            y = h - 40
            c.setFont("Helvetica-Bold", 9)
            c.drawString(36, y, "Fecha/Hora")
            c.drawString(130, y, "Titulo")
            c.drawString(305, y, "Grupo")
            c.drawString(390, y, "Lugar")
            c.drawString(470, y, "Tipo")
            y -= 12
            c.setFont("Helvetica", 8)

        fecha_hora = f"{fila['fecha_evento']} {fila['hora_evento'] or ''}".strip()
        titulo = (fila["titulo"] or "")[:34]
        grupo = (fila["grupo"] or "-")[:14]
        lugar = (fila["lugar"] or "-")[:14]
        tipo = (fila["tipo"] or "-")[:10]
        c.drawString(36, y, fecha_hora)
        c.drawString(130, y, titulo)
        c.drawString(305, y, grupo)
        c.drawString(390, y, lugar)
        c.drawString(470, y, tipo)
        y -= 11

    if not eventos:
        c.drawString(36, y, "No hay eventos para este mes con los filtros actuales.")

    c.showPage()
    c.save()
    buff.seek(0)
    nombre = f"calendario_{anio:04d}_{mes:02d}.pdf"
    return Response(
        buff.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


def _parse_import_rows_from_upload(file_storage):
    filename = (getattr(file_storage, "filename", "") or "").lower()
    raw = file_storage.read()
    if not raw:
        return []

    rows = []
    if filename.endswith(".xlsx"):
        try:
            import openpyxl
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError(f"Para importar Excel instala openpyxl ({exc}).")
        from io import BytesIO as _B

        wb = openpyxl.load_workbook(_B(raw), data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return []
        headers = [str(c or "").strip().lower() for c in data[0]]
        for line in data[1:]:
            row_map = {headers[i]: (line[i] if i < len(line) else "") for i in range(len(headers))}
            rows.append(row_map)
        return rows

    text = raw.decode("utf-8", errors="ignore")
    sample = text.splitlines()[0] if text.splitlines() else ""
    delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(StringIO(text), delimiter=delim)
    for r in reader:
        row = {}
        for k, v in (r or {}).items():
            row[(k or "").strip().lower()] = (v or "").strip()
        rows.append(row)
    return rows


def _normalizar_import_event_row(row):
    def pick(*keys):
        for k in keys:
            v = (row.get(k) or "").strip() if isinstance(row, dict) else ""
            if v:
                return v
        return ""

    titulo = pick("titulo", "title")
    fecha = pick("fecha", "fecha_evento", "date")
    hora = pick("hora", "hora_evento", "time")
    grupo = pick("grupo", "group")
    lugar = pick("lugar", "location")
    tipo = pick("tipo", "type")
    descripcion = pick("descripcion", "description")
    fecha_obj = parsear_fecha_limite(fecha)
    if not titulo or not fecha_obj:
        return None
    hora_final = ""
    if hora and re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", hora):
        hora_final = hora
    return {
        "titulo": titulo[:120],
        "fecha_evento": fecha_obj.strftime("%Y-%m-%d"),
        "hora_evento": hora_final,
        "grupo": grupo[:60],
        "lugar": lugar[:60],
        "tipo": tipo[:40],
        "descripcion": descripcion[:240],
    }


def importar_eventos():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))
    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    preview = session.get("cal_import_preview") or []
    if request.method == "POST":
        archivo = request.files.get("archivo")
        if not archivo or not archivo.filename:
            flash("Selecciona un archivo CSV o XLSX.", "warning")
            return redirect(url_for("calendario_bp.importar_eventos"))
        try:
            raw_rows = _parse_import_rows_from_upload(archivo)
        except Exception as exc:  # noqa: BLE001
            flash(f"No se pudo leer archivo: {exc}", "danger")
            return redirect(url_for("calendario_bp.importar_eventos"))
        preview = []
        for r in raw_rows:
            n = _normalizar_import_event_row(r)
            if n:
                preview.append(n)
        session["cal_import_preview"] = preview[:500]
        flash(f"Vista previa cargada: {len(preview)} evento(s) validos.", "info")
        return redirect(url_for("calendario_bp.importar_eventos"))

    return render_template("calendario/importar.html", preview=preview, total_preview=len(preview))


def confirmar_importacion_eventos():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))
    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    preview = session.get("cal_import_preview") or []
    if not preview:
        flash("No hay vista previa para importar.", "warning")
        return redirect(url_for("calendario_bp.importar_eventos"))

    conn = get_connection()
    cursor = conn.cursor()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total = 0
    for ev in preview:
        cursor.execute(
            """
            INSERT INTO calendario_eventos
            (user_id, titulo, fecha_evento, hora_evento, grupo, lugar, tipo, descripcion, creado_en)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                usuario_id,
                ev["titulo"],
                ev["fecha_evento"],
                ev["hora_evento"] or None,
                ev["grupo"] or None,
                ev["lugar"] or None,
                ev["tipo"] or None,
                ev["descripcion"] or None,
                now,
            ),
        )
        total += 1
    conn.commit()
    conn.close()
    session.pop("cal_import_preview", None)
    registrar_auditoria(
        usuario_id,
        "calendario",
        "importar",
        "evento",
        None,
        {"total": total},
    )
    flash(f"Importacion completada: {total} evento(s).", "success")
    return redirect(url_for("calendario_bp.calendario"))


def _carpeta_mensajes_solicitada():
    return mensajes_logic.carpeta_mensajes_solicitada(request)


def _buscar_mensajes_usuario(usuario_id, carpeta, q, limit=10, offset=0):
    return mensajes_logic.buscar_mensajes_usuario(get_connection, usuario_id, carpeta, q, limit, offset)


def _contar_mensajes_filtrados(usuario_id, carpeta, q):
    return mensajes_logic.contar_mensajes_filtrados(get_connection, usuario_id, carpeta, q)


def _conteo_mensajes(usuario_id):
    return mensajes_logic.conteo_mensajes(get_connection, usuario_id)


def mensajes():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    carpeta = _carpeta_mensajes_solicitada()
    q = request.args.get("q", "").strip()
    try:
        page = max(1, int(request.args.get("page", "1")))
    except ValueError:
        page = 1
    page_size = 10
    total = _contar_mensajes_filtrados(usuario_id, carpeta, q)
    total_pages = max(1, (total + page_size - 1) // page_size)
    if page > total_pages:
        page = total_pages
    offset = (page - 1) * page_size

    lista = _buscar_mensajes_usuario(usuario_id, carpeta, q, page_size, offset)
    conteo = _conteo_mensajes(usuario_id)

    if carpeta == "entrada":
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE mensajes SET leido = 1 WHERE user_id = ? AND carpeta = 'entrada' AND leido = 0",
            (usuario_id,),
        )
        conn.commit()
        conn.close()

    return render_template(
        "mensajes/mensajes.html",
        carpeta=carpeta,
        q=q,
        mensajes=lista,
        conteo=conteo,
        page=page,
        total_pages=total_pages,
        total_mensajes=total,
    )


def nuevo_mensaje():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    remitente = (session.get("usuario") or "").strip()
    if usuario_id is None or not remitente:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    destinatario = request.form.get("destinatario", "").strip()
    asunto = request.form.get("asunto", "").strip()
    cuerpo = request.form.get("cuerpo", "").strip()

    if not destinatario or not asunto or not cuerpo:
        flash("Completa destinatario, asunto y mensaje.", "warning")
        return redirect(url_for("mensajes_bp.mensajes", carpeta="enviados"))

    marca = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT id, username FROM usuarios WHERE lower(username) = ?", (destinatario.lower(),))
    destino = cursor.fetchone()
    if destino is None:
        conn.close()
        flash("El destinatario no existe.", "danger")
        return redirect(url_for("mensajes_bp.mensajes", carpeta="enviados"))

    cursor.execute(
        """
        INSERT INTO mensajes (user_id, carpeta, remitente, destinatario, asunto, cuerpo, leido, creado_en)
        VALUES (?, 'enviados', ?, ?, ?, ?, 1, ?)
        """,
        (usuario_id, remitente, destino["username"], asunto, cuerpo, marca),
    )
    cursor.execute(
        """
        INSERT INTO mensajes (user_id, carpeta, remitente, destinatario, asunto, cuerpo, leido, creado_en)
        VALUES (?, 'entrada', ?, ?, ?, ?, 0, ?)
        """,
        (destino["id"], remitente, destino["username"], asunto, cuerpo, marca),
    )

    conn.commit()
    mensaje_id = cursor.lastrowid
    conn.close()
    registrar_auditoria(
        usuario_id,
        "mensajes",
        "crear",
        "mensaje",
        mensaje_id,
        {"destinatario": destino["username"], "asunto": asunto},
    )
    flash("Mensaje enviado.", "success")
    return redirect(url_for("mensajes_bp.mensajes", carpeta="enviados"))


def mover_mensaje(mensaje_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    accion = request.form.get("accion", "").strip().lower()
    carpeta_actual = request.form.get("carpeta", "entrada").strip().lower()
    if carpeta_actual not in ("entrada", "enviados", "papelera"):
        carpeta_actual = "entrada"

    destino = "papelera" if accion == "papelera" else "entrada"
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE mensajes SET carpeta = ? WHERE id = ? AND user_id = ?",
        (destino, mensaje_id, usuario_id),
    )
    conn.commit()
    actualizado = cursor.rowcount
    conn.close()

    if actualizado:
        if destino == "papelera":
            registrar_auditoria(
                usuario_id,
                "mensajes",
                "mover",
                "mensaje",
                mensaje_id,
                {"destino": "papelera"},
            )
            flash("Mensaje movido a papelera.", "info")
            return redirect(url_for("mensajes_bp.mensajes", carpeta="papelera"))
        registrar_auditoria(
            usuario_id,
            "mensajes",
            "mover",
            "mensaje",
            mensaje_id,
            {"destino": "entrada"},
        )
        flash("Mensaje restaurado a entrada.", "success")
        return redirect(url_for("mensajes_bp.mensajes", carpeta="entrada"))

    flash("No se pudo mover el mensaje.", "warning")
    return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual))


def editar_mensaje(mensaje_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    carpeta_actual = request.form.get("carpeta", "entrada").strip().lower()
    if carpeta_actual not in ("entrada", "enviados", "papelera"):
        carpeta_actual = "entrada"
    q = request.form.get("q", "").strip()

    asunto = request.form.get("asunto", "").strip()
    cuerpo = request.form.get("cuerpo", "").strip()
    if not asunto or not cuerpo:
        flash("Asunto y mensaje son obligatorios para editar.", "warning")
        return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual, q=q))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT carpeta FROM mensajes WHERE id = ? AND user_id = ?",
        (mensaje_id, usuario_id),
    )
    fila = cursor.fetchone()
    if fila is None:
        conn.close()
        flash("No se pudo editar el mensaje.", "warning")
        return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual, q=q))

    if (fila["carpeta"] or "").strip().lower() != "enviados":
        conn.close()
        flash("Solo puedes editar mensajes en la bandeja Enviados.", "warning")
        return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual, q=q))

    cursor.execute(
        "UPDATE mensajes SET asunto = ?, cuerpo = ? WHERE id = ? AND user_id = ?",
        (asunto, cuerpo, mensaje_id, usuario_id),
    )
    conn.commit()
    actualizado = cursor.rowcount
    conn.close()

    if actualizado:
        registrar_auditoria(
            usuario_id,
            "mensajes",
            "editar",
            "mensaje",
            mensaje_id,
            {"asunto": asunto},
        )
        flash("Mensaje actualizado.", "success")
    else:
        flash("No se pudo editar el mensaje.", "warning")

    return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual, q=q))


def eliminar_mensaje(mensaje_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    carpeta_actual = request.form.get("carpeta", "papelera").strip().lower()
    if carpeta_actual not in ("entrada", "enviados", "papelera"):
        carpeta_actual = "papelera"
    q = request.form.get("q", "").strip()

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "DELETE FROM mensajes WHERE id = ? AND user_id = ? AND carpeta = 'papelera'",
        (mensaje_id, usuario_id),
    )
    conn.commit()
    eliminado = cursor.rowcount
    conn.close()

    if eliminado:
        registrar_auditoria(usuario_id, "mensajes", "eliminar", "mensaje", mensaje_id, {})
        flash("Mensaje eliminado definitivamente.", "info")
    else:
        flash("Solo puedes eliminar mensajes que esten en papelera.", "warning")
    return redirect(url_for("mensajes_bp.mensajes", carpeta=carpeta_actual, q=q))


def inicio_metas():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))
    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, meta, monto, ahorrado, fecha_limite
        FROM metas
        WHERE user_id = ?
        ORDER BY id DESC
        LIMIT 6
        """,
        (usuario_id,),
    )
    metas_recientes = cursor.fetchall()
    cursor.execute(
        """
        SELECT id, titulo, fecha_evento, hora_evento
        FROM calendario_eventos
        WHERE user_id = ? AND date(fecha_evento) >= date('now')
        ORDER BY fecha_evento ASC, coalesce(hora_evento, '99:99') ASC, id ASC
        LIMIT 6
        """,
        (usuario_id,),
    )
    eventos_proximos = cursor.fetchall()
    cursor.execute(
        """
        SELECT asunto, remitente, creado_en
        FROM mensajes
        WHERE user_id = ? AND carpeta = 'entrada'
        ORDER BY datetime(creado_en) DESC, id DESC
        LIMIT 5
        """,
        (usuario_id,),
    )
    mensajes_recientes = cursor.fetchall()
    conn.close()

    actividad_reciente = obtener_actividad_reciente_usuario(usuario_id, limite=10)
    kpis = obtener_kpis_usuario(usuario_id)
    return render_template(
        "metas/inicio.html",
        kpis=kpis,
        actividad_reciente=actividad_reciente,
        metas_recientes=metas_recientes,
        eventos_proximos=eventos_proximos,
        mensajes_recientes=mensajes_recientes,
    )


@app.get("/buscar")
def buscar_global():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))
    rate = aplicar_rate_limit_api("api_busqueda")
    if rate:
        return rate
    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    q = (request.args.get("q") or "").strip()
    resultados = {"metas": [], "eventos": [], "mensajes": []}
    if q:
        like = f"%{q}%"
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT id, meta, fecha_limite
            FROM metas
            WHERE user_id = ? AND (meta LIKE ? OR fecha_limite LIKE ?)
            ORDER BY id DESC
            LIMIT 20
            """,
            (usuario_id, like, like),
        )
        resultados["metas"] = cursor.fetchall()
        cursor.execute(
            """
            SELECT id, titulo, fecha_evento, coalesce(grupo,'') AS grupo, coalesce(lugar,'') AS lugar
            FROM calendario_eventos
            WHERE user_id = ? AND (titulo LIKE ? OR coalesce(grupo,'') LIKE ? OR coalesce(lugar,'') LIKE ?)
            ORDER BY fecha_evento DESC, id DESC
            LIMIT 20
            """,
            (usuario_id, like, like, like),
        )
        resultados["eventos"] = cursor.fetchall()
        cursor.execute(
            """
            SELECT id, carpeta, coalesce(asunto,'(Sin asunto)') AS asunto, coalesce(remitente,'-') AS remitente, creado_en
            FROM mensajes
            WHERE user_id = ? AND (coalesce(asunto,'') LIKE ? OR coalesce(cuerpo,'') LIKE ? OR coalesce(remitente,'') LIKE ?)
            ORDER BY datetime(creado_en) DESC, id DESC
            LIMIT 25
            """,
            (usuario_id, like, like, like),
        )
        resultados["mensajes"] = cursor.fetchall()
        conn.close()
    return render_template("buscar.html", q=q, resultados=resultados)


def crear_meta():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    meta = request.form["meta"].strip()
    fecha_limite = request.form["fecha_limite"]
    fecha_obj = parsear_fecha_limite(fecha_limite)

    try:
        monto = float(request.form["monto"])
        ahorrado = float(request.form["ahorrado"])
    except ValueError:
        flash("Monto y ahorrado deben ser numericos.", "danger")
        return redirect(url_for("metas_bp.inicio_metas"))

    if fecha_obj is None:
        flash("La fecha limite no tiene un formato valido.", "danger")
        return redirect(url_for("metas_bp.inicio_metas"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO metas (meta, monto, ahorrado, fecha_limite, user_id)
        VALUES (?, ?, ?, ?, ?)
        """,
        (meta, monto, ahorrado, fecha_obj.strftime("%Y-%m-%d"), usuario_id),
    )
    meta_id = cursor.lastrowid
    conn.commit()
    conn.close()
    registrar_auditoria(
        usuario_id,
        "metas",
        "crear",
        "meta",
        meta_id,
        {"meta": meta, "fecha_limite": fecha_obj.strftime("%Y-%m-%d")},
    )

    flash("Meta creada correctamente.", "success")
    return redirect(url_for("metas_bp.inicio_metas"))


def ver_metas():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM metas WHERE user_id = ? ORDER BY id DESC", (usuario_id,))
    datos = cursor.fetchall()
    conn.close()

    metas = metas_logic.construir_metas_para_vista(datos, parsear_fecha_limite)

    return render_template("metas/ver_metas.html", metas=metas)


def editar(meta_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM metas WHERE id = ? AND user_id = ?", (meta_id, usuario_id))
    meta = cursor.fetchone()
    conn.close()

    if meta is None:
        flash("Meta no encontrada.", "warning")
        return redirect(url_for("metas_bp.ver_metas"))

    fecha_para_input = parsear_fecha_limite(meta["fecha_limite"])
    fecha_para_input = fecha_para_input.strftime("%Y-%m-%d") if fecha_para_input else ""
    return render_template("metas/editar.html", meta=meta, fecha_limite_input=fecha_para_input)


def actualizar(meta_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    meta = request.form["meta"].strip()
    fecha_limite = request.form["fecha_limite"]
    fecha_obj = parsear_fecha_limite(fecha_limite)

    try:
        monto = float(request.form["monto"])
        ahorrado = float(request.form["ahorrado"])
    except ValueError:
        flash("Monto y ahorrado deben ser numericos.", "danger")
        return redirect(url_for("metas_bp.editar", meta_id=meta_id))

    if fecha_obj is None:
        flash("La fecha limite no tiene un formato valido.", "danger")
        return redirect(url_for("metas_bp.editar", meta_id=meta_id))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        UPDATE metas
        SET meta = ?, monto = ?, ahorrado = ?, fecha_limite = ?
        WHERE id = ? AND user_id = ?
        """,
        (meta, monto, ahorrado, fecha_obj.strftime("%Y-%m-%d"), meta_id, usuario_id),
    )
    conn.commit()
    actualizadas = cursor.rowcount
    conn.close()

    if actualizadas == 0:
        flash("No tienes permisos para actualizar esa meta.", "warning")
        return redirect(url_for("metas_bp.ver_metas"))

    registrar_auditoria(
        usuario_id,
        "metas",
        "editar",
        "meta",
        meta_id,
        {"meta": meta, "fecha_limite": fecha_obj.strftime("%Y-%m-%d")},
    )
    flash("Meta actualizada.", "success")
    return redirect(url_for("metas_bp.ver_metas"))


def eliminar(meta_id):
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))

    usuario_id = obtener_usuario_id_actual()
    if usuario_id is None:
        flash("Tu sesion no es valida. Vuelve a iniciar sesion.", "warning")
        return redirect(url_for("auth_bp.logout"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM metas WHERE id = ? AND user_id = ?", (meta_id, usuario_id))
    conn.commit()
    eliminadas = cursor.rowcount
    conn.close()

    if eliminadas == 0:
        flash("No tienes permisos para eliminar esa meta.", "warning")
        return redirect(url_for("metas_bp.ver_metas"))

    registrar_auditoria(usuario_id, "metas", "eliminar", "meta", meta_id, {})
    flash("Meta eliminada.", "info")
    return redirect(url_for("metas_bp.ver_metas"))


def _resumen_health(include_private=False):
    estado = {"ok": True, "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    checks = {}

    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT 1")
        cursor.fetchone()
        conn.close()
        checks["db"] = "ok"
    except Exception as exc:  # noqa: BLE001
        checks["db"] = f"error: {exc}"
        estado["ok"] = False

    checks["mail"] = "ok" if proveedor_correo_configurado() else "not_configured"
    checks["sms"] = "ok" if sms_configurado() else "not_configured"
    checks["push"] = "ok" if push_web_configurado() else "not_configured"

    estado["checks"] = checks
    if include_private:
        estado["env"] = {
            "db_path": DB_PATH,
            "backup_dir": BACKUP_DIR,
            "mail_route": nombre_ruta_proveedores(),
        }
    return estado


def _tail_log(ruta_relativa, lineas=30):
    ruta = os.path.join(BASE_DIR, ruta_relativa)
    if not os.path.exists(ruta):
        return []
    try:
        with open(ruta, "r", encoding="utf-8", errors="ignore") as fh:
            data = fh.readlines()
        return [ln.rstrip("\n") for ln in data[-lineas:]]
    except OSError:
        return []


def _resumen_reporte_diario():
    ruta = os.path.join(BASE_DIR, "logs", "reporte_diario_last.json")
    if not os.path.exists(ruta):
        return {
            "status": "red",
            "label": "Sin reporte",
            "summary": "No existe logs/reporte_diario_last.json",
            "timestamp": "-",
            "items": [],
        }

    try:
        with open(ruta, "r", encoding="utf-8", errors="ignore") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        return {
            "status": "red",
            "label": "Reporte invalido",
            "summary": f"No se pudo leer el reporte: {exc}",
            "timestamp": "-",
            "items": [],
        }

    ts = str(data.get("timestamp") or "-")
    overall_ok = bool(data.get("overall_ok"))
    critical_ok = bool(data.get("critical_ok"))
    tasks = data.get("tasks") or []
    fail_tasks = sum(1 for t in tasks if not bool(t.get("ok")))
    logs = data.get("logs") or []
    warn_logs = sum(1 for l in logs if not bool(l.get("ok")))

    status = "green" if overall_ok else ("yellow" if critical_ok else "red")
    label = "OK" if status == "green" else ("Atencion" if status == "yellow" else "Alerta")
    summary = (
        f"Tareas fail: {fail_tasks} | Logs con alertas: {warn_logs} | "
        f"Notificaciones hoy: {int((data.get('metrics') or {}).get('runs_today', 0) or 0)} ejecucion(es)"
    )

    if ts != "-":
        try:
            dt = datetime.strptime(ts, "%Y-%m-%d %H:%M:%S")
            horas = (datetime.now() - dt).total_seconds() / 3600.0
            if horas > 30:
                status = "red"
                label = "Alerta"
                summary += f" | Reporte desactualizado ({horas:.1f}h)"
            elif horas > 24 and status == "green":
                status = "yellow"
                label = "Atencion"
                summary += f" | Reporte algo antiguo ({horas:.1f}h)"
        except ValueError:
            pass

    items = [
        {"name": "Estado general", "value": "OK" if overall_ok else "ALERTA"},
        {"name": "Critico", "value": "OK" if critical_ok else "FAIL"},
        {"name": "Tareas con fallo", "value": str(fail_tasks)},
        {"name": "Logs con alerta", "value": str(warn_logs)},
    ]
    return {"status": status, "label": label, "summary": summary, "timestamp": ts, "items": items}


def _historial_reporte_diario(limit=7):
    path = os.path.join(BASE_DIR, "logs", "reporte_diario_history.jsonl")
    if not os.path.exists(path):
        return []
    rows = []
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            for raw in fh:
                line = raw.strip()
                if not line:
                    continue
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    except OSError:
        return []

    rows = rows[-max(1, int(limit)) :]
    points = []
    for item in rows:
        status = str(item.get("status") or "red")
        score = 100 if status == "green" else (60 if status == "yellow" else 20)
        ts = str(item.get("timestamp") or "")
        label = ts[5:10] if len(ts) >= 10 else ts
        points.append({"label": label, "status": status, "score": score})
    return points


def _estado_tarea_windows(nombre_tarea):
    try:
        proc = subprocess.run(
            ["schtasks", "/Query", "/TN", nombre_tarea, "/V", "/FO", "LIST"],
            capture_output=True,
            text=True,
            timeout=12,
            check=False,
        )
    except (FileNotFoundError, OSError):
        return {"task": nombre_tarea, "ok": False, "summary": "schtasks no disponible."}

    if proc.returncode != 0:
        return {
            "task": nombre_tarea,
            "ok": False,
            "summary": (proc.stderr or proc.stdout or "No disponible").strip(),
        }

    texto = proc.stdout or ""
    campos = {}
    for raw in texto.splitlines():
        if ":" not in raw:
            continue
        k, v = raw.split(":", 1)
        campos[k.strip()] = v.strip()

    estado = campos.get("Estado", "-")
    proxima = campos.get("Hora próxima ejecución", "-")
    ultimo_resultado = campos.get("Último resultado", "-")
    ejecuta_como = campos.get("Ejecutar como usuario", "-")
    summary = f"Estado={estado} | Proxima={proxima} | Ultimo={ultimo_resultado} | Usuario={ejecuta_como}"
    return {"task": nombre_tarea, "ok": True, "summary": summary}


def _admin_guard():
    if not usuario_autenticado():
        return redirect(url_for("auth_bp.login"))
    if not usuario_es_admin():
        abort(403)
    return None


def _admin_rate_guard():
    rate = aplicar_rate_limit_api("api_admin")
    if rate:
        return rate
    return None


def _filtros_auditoria_desde_request():
    return {
        "fecha_desde": (request.args.get("fecha_desde") or "").strip(),
        "fecha_hasta": (request.args.get("fecha_hasta") or "").strip(),
        "actor": (request.args.get("actor") or "").strip(),
        "modulo": (request.args.get("modulo") or "").strip(),
        "accion": (request.args.get("accion") or "").strip(),
        "detalle_q": (request.args.get("detalle_q") or "").strip(),
    }


def _where_params_auditoria(filtros):
    where = ["1=1"]
    params = []

    fecha_desde = filtros.get("fecha_desde", "")
    fecha_hasta = filtros.get("fecha_hasta", "")
    actor = filtros.get("actor", "")
    modulo = filtros.get("modulo", "")
    accion = filtros.get("accion", "")
    detalle_q = filtros.get("detalle_q", "")

    if fecha_desde:
        where.append("date(a.creado_en) >= date(?)")
        params.append(fecha_desde)
    if fecha_hasta:
        where.append("date(a.creado_en) <= date(?)")
        params.append(fecha_hasta)
    if actor:
        where.append("lower(u.username) = lower(?)")
        params.append(actor)
    if modulo:
        where.append("lower(a.modulo) = lower(?)")
        params.append(modulo)
    if accion:
        where.append("lower(a.accion) = lower(?)")
        params.append(accion)
    if detalle_q:
        where.append("lower(coalesce(a.detalle,'')) LIKE lower(?)")
        params.append(f"%{detalle_q}%")
    return where, params


def _query_auditoria(cursor, filtros, limit=40, offset=0):
    where, params = _where_params_auditoria(filtros)
    sql = f"""
        SELECT a.id, a.modulo, a.accion, a.entidad, a.entidad_id, a.detalle, a.creado_en, u.username AS actor
        FROM audit_log a
        LEFT JOIN usuarios u ON u.id = a.actor_user_id
        WHERE {' AND '.join(where)}
        ORDER BY datetime(a.creado_en) DESC, a.id DESC
    """
    if limit and int(limit) > 0:
        sql += " LIMIT ? OFFSET ?"
        params.append(int(limit))
        params.append(max(0, int(offset or 0)))

    cursor.execute(sql, tuple(params))
    return cursor.fetchall()


def _contar_auditoria(cursor, filtros):
    where, params = _where_params_auditoria(filtros)
    sql = f"""
        SELECT COUNT(*)
        FROM audit_log a
        LEFT JOIN usuarios u ON u.id = a.actor_user_id
        WHERE {' AND '.join(where)}
    """
    cursor.execute(sql, tuple(params))
    fila = cursor.fetchone()
    return int(fila[0] if fila else 0)


def _preparar_filas_auditoria(rows):
    prepared = []
    for r in rows:
        item = dict(r)
        raw = str(item.get("detalle") or "").strip()
        parsed = None
        if raw.startswith("{") and raw.endswith("}"):
            try:
                parsed = json.loads(raw)
            except json.JSONDecodeError:
                parsed = None
        if isinstance(parsed, dict):
            visible = {k: v for k, v in parsed.items() if k not in ("ip", "user_agent")}
            ip = str(parsed.get("ip") or "")
            ua = str(parsed.get("user_agent") or "")
            if visible:
                partes = []
                for k, v in visible.items():
                    label = str(k).replace("_", " ").strip().capitalize()
                    partes.append(f"{label}: {v}")
                item["detalle_resumen"] = " | ".join(partes)
            else:
                item["detalle_resumen"] = "-"
            item["detalle_ip"] = ip
            item["detalle_ua"] = ua
            item["detalle_raw"] = raw
        else:
            item["detalle_resumen"] = raw or "-"
            item["detalle_ip"] = ""
            item["detalle_ua"] = ""
            item["detalle_raw"] = raw
        prepared.append(item)
    return prepared


def _saved_audit_filters(user_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, nombre, fecha_desde, fecha_hasta, actor, modulo, accion, detalle_q, creado_en
        FROM saved_audit_filters
        WHERE user_id = ?
        ORDER BY datetime(creado_en) DESC, id DESC
        LIMIT 20
        """,
        (user_id,),
    )
    rows = cursor.fetchall()
    conn.close()
    return rows


def _metricas_operativas_7d():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT date(creado_en) AS d, modulo, COUNT(*) AS total
        FROM audit_log
        WHERE date(creado_en) >= date('now', '-6 day')
          AND accion = 'crear'
          AND modulo IN ('metas', 'calendario', 'mensajes')
        GROUP BY date(creado_en), modulo
        ORDER BY d ASC
        """
    )
    rows = cursor.fetchall()
    cursor.execute(
        """
        SELECT date(creado_en) AS d, COUNT(*) AS total
        FROM auth_login_attempts
        WHERE date(creado_en) >= date('now', '-6 day') AND success = 0
        GROUP BY date(creado_en)
        ORDER BY d ASC
        """
    )
    fail_rows = cursor.fetchall()
    conn.close()

    days = []
    for i in range(6, -1, -1):
        d = (datetime.now().date() - timedelta(days=i)).strftime("%Y-%m-%d")
        days.append(d)
    data = {d: {"metas": 0, "calendario": 0, "mensajes": 0, "login_fail": 0} for d in days}
    for r in rows:
        day = r["d"]
        mod = r["modulo"]
        if day in data and mod in data[day]:
            data[day][mod] = int(r["total"] or 0)
    for r in fail_rows:
        day = r["d"]
        if day in data:
            data[day]["login_fail"] = int(r["total"] or 0)

    serie = []
    totals = {"metas": 0, "calendario": 0, "mensajes": 0, "login_fail": 0}
    for d in days:
        item = {"dia": d[5:], **data[d]}
        for k in totals:
            totals[k] += item[k]
        serie.append(item)
    return {"serie": serie, "totals": totals}


@app.get("/admin/dashboard")
def admin_dashboard():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    admin_user_id = obtener_usuario_id_actual()
    conn = get_connection()
    cursor = conn.cursor()
    filtros = _filtros_auditoria_desde_request()
    try:
        page = max(1, int(request.args.get("page", "1")))
    except ValueError:
        page = 1
    page_size = 25
    total_audit = _contar_auditoria(cursor, filtros)
    total_pages = max(1, (total_audit + page_size - 1) // page_size)
    if page > total_pages:
        page = total_pages
    offset = (page - 1) * page_size
    audit_rows = _preparar_filas_auditoria(_query_auditoria(cursor, filtros, limit=page_size, offset=offset))
    cursor.execute("SELECT username FROM usuarios ORDER BY lower(username) ASC")
    actores_disponibles = [r["username"] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT modulo FROM audit_log ORDER BY lower(modulo) ASC")
    modulos_disponibles = [r["modulo"] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT accion FROM audit_log ORDER BY lower(accion) ASC")
    acciones_disponibles = [r["accion"] for r in cursor.fetchall()]
    conn.close()
    saved_filters = _saved_audit_filters(admin_user_id) if admin_user_id else []

    tareas = [
        _estado_tarea_windows(r"\Meta90_Notificaciones_Manana"),
        _estado_tarea_windows(r"\Meta90_Notificaciones_Noche"),
        _estado_tarea_windows(r"\Meta90_DB_Backup"),
        _estado_tarea_windows(r"\Meta90_DB_Restore_Verify"),
        _estado_tarea_windows(r"\Meta90_Task_Monitor"),
        _estado_tarea_windows(r"\Meta90_Reporte_Diario"),
        _estado_tarea_windows(r"\Meta90_Log_Rotation"),
        _estado_tarea_windows(r"\Meta90_Semaforo_Alert"),
    ]

    logs = {
        "notificaciones": _tail_log(os.path.join("logs", "notificaciones_scheduler.log"), 20),
        "backup": _tail_log(os.path.join("logs", "db_backup.log"), 20),
        "restore_verify": _tail_log(os.path.join("logs", "db_restore_verify.log"), 20),
        "task_monitor": _tail_log(os.path.join("logs", "task_monitor.log"), 20),
        "log_rotation": _tail_log(os.path.join("logs", "log_rotation.log"), 20),
        "semaforo_alerta": _tail_log(os.path.join("logs", "semaforo_alerta.log"), 20),
    }
    reporte_diario = _resumen_reporte_diario()
    historial_reporte = _historial_reporte_diario(limit=7)
    mantenimiento_notificaciones = notificaciones_en_mantenimiento()
    metricas_7d = _metricas_operativas_7d()

    return render_template(
        "admin/dashboard.html",
        health=_resumen_health(include_private=True),
        tareas=tareas,
        logs=logs,
        reporte_diario=reporte_diario,
        historial_reporte=historial_reporte,
        mantenimiento_notificaciones=mantenimiento_notificaciones,
        metricas_7d=metricas_7d,
        saved_filters=saved_filters,
        audit_rows=audit_rows,
        filtros=filtros,
        page=page,
        total_pages=total_pages,
        total_audit=total_audit,
        actores_disponibles=actores_disponibles,
        modulos_disponibles=modulos_disponibles,
        acciones_disponibles=acciones_disponibles,
    )


@app.get("/admin/reporte-diario/status")
def admin_reporte_diario_status():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    return jsonify(_resumen_reporte_diario())


@app.post("/admin/audit/filtros/guardar")
def admin_guardar_filtro_auditoria():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    user_id = obtener_usuario_id_actual()
    if not user_id:
        return redirect(url_for("auth_bp.login"))
    nombre = (request.form.get("nombre") or "").strip()
    if not nombre:
        flash("Nombre de filtro requerido.", "warning")
        return redirect(url_for("admin_dashboard"))
    filtros = {
        "fecha_desde": (request.form.get("fecha_desde") or "").strip(),
        "fecha_hasta": (request.form.get("fecha_hasta") or "").strip(),
        "actor": (request.form.get("actor") or "").strip(),
        "modulo": (request.form.get("modulo") or "").strip(),
        "accion": (request.form.get("accion") or "").strip(),
        "detalle_q": (request.form.get("detalle_q") or "").strip(),
    }
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO saved_audit_filters
        (user_id, nombre, fecha_desde, fecha_hasta, actor, modulo, accion, detalle_q, creado_en)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user_id,
            nombre[:60],
            filtros["fecha_desde"],
            filtros["fecha_hasta"],
            filtros["actor"],
            filtros["modulo"],
            filtros["accion"],
            filtros["detalle_q"],
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    )
    conn.commit()
    conn.close()
    flash("Filtro guardado.", "success")
    return redirect(url_for("admin_dashboard", **filtros))


@app.post("/admin/audit/filtros/<int:filter_id>/eliminar")
def admin_eliminar_filtro_auditoria(filter_id):
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    user_id = obtener_usuario_id_actual()
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM saved_audit_filters WHERE id = ? AND user_id = ?", (filter_id, user_id))
    conn.commit()
    ok = cursor.rowcount > 0
    conn.close()
    flash("Filtro eliminado." if ok else "No se pudo eliminar filtro.", "info" if ok else "warning")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/notificaciones/mantenimiento")
def admin_toggle_mantenimiento_notificaciones():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    accion = (request.form.get("accion") or "").strip().lower()
    activar = accion == "on"
    set_mantenimiento_notificaciones(activar, obtener_usuario_id_actual())
    flash(
        "Mantenimiento de notificaciones activado." if activar else "Mantenimiento de notificaciones desactivado.",
        "warning" if activar else "success",
    )
    return redirect(url_for("admin_dashboard"))


@app.get("/admin/audit/export.csv")
def admin_audit_export_csv():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    filtros = _filtros_auditoria_desde_request()
    conn = get_connection()
    cursor = conn.cursor()
    audit_rows = _query_auditoria(cursor, filtros, limit=0)
    conn.close()

    salida = StringIO()
    writer = csv.writer(salida)
    writer.writerow(["id", "creado_en", "actor", "modulo", "accion", "entidad", "entidad_id", "detalle"])
    for a in audit_rows:
        writer.writerow(
            [
                a["id"],
                a["creado_en"],
                a["actor"] or "",
                a["modulo"],
                a["accion"],
                a["entidad"],
                a["entidad_id"] or "",
                a["detalle"] or "",
            ]
        )

    contenido = salida.getvalue()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = f"audit_log_{ts}.csv"
    return Response(
        contenido,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


@app.get("/admin/usuarios")
def admin_usuarios():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, username, email, telefono, coalesce(rol, 'editor') AS rol
        FROM usuarios
        ORDER BY lower(username) ASC
        """
    )
    usuarios = cursor.fetchall()
    permisos = {}
    for u in usuarios:
        user_id = int(u["id"])
        permisos[user_id] = {}
        for modulo in MODULES_WITH_PERMS:
            p = obtener_permiso_modulo_usuario(user_id, modulo)
            permisos[user_id][modulo] = p
    conn.close()
    return render_template("admin/usuarios.html", usuarios=usuarios, permisos=permisos, modules_with_perms=MODULES_WITH_PERMS)


@app.post("/admin/usuarios/<int:user_id>/rol")
def admin_actualizar_rol(user_id):
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    nuevo_rol = (request.form.get("rol") or "").strip().lower()
    if nuevo_rol not in ("admin", "editor", "lector"):
        flash("Rol invalido.", "warning")
        return redirect(url_for("admin_usuarios"))

    actor_id = obtener_usuario_id_actual()
    if actor_id == user_id and nuevo_rol != "admin":
        flash("No puedes quitarte el rol admin a ti mismo.", "warning")
        return redirect(url_for("admin_usuarios"))

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE usuarios SET rol = ? WHERE id = ?", (nuevo_rol, user_id))
    conn.commit()
    ok = cursor.rowcount > 0
    conn.close()
    if ok:
        registrar_auditoria(
            actor_id,
            "admin",
            "cambiar_rol",
            "usuario",
            user_id,
            {"rol": nuevo_rol},
        )
        flash("Rol actualizado.", "success")
    else:
        flash("Usuario no encontrado.", "warning")
    return redirect(url_for("admin_usuarios"))


@app.post("/admin/usuarios/<int:user_id>/permiso")
def admin_actualizar_permiso_modulo(user_id):
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    modulo = (request.form.get("modulo") or "").strip().lower()
    if modulo not in MODULES_WITH_PERMS:
        flash("Modulo invalido.", "warning")
        return redirect(url_for("admin_usuarios"))
    can_view = 1 if request.form.get("can_view") == "1" else 0
    can_edit = 1 if request.form.get("can_edit") == "1" else 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO user_module_permissions (user_id, modulo, can_view, can_edit, actualizado_en)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(user_id, modulo) DO UPDATE SET
            can_view = excluded.can_view,
            can_edit = excluded.can_edit,
            actualizado_en = excluded.actualizado_en
        """,
        (user_id, modulo, can_view, can_edit, now),
    )
    conn.commit()
    conn.close()
    registrar_auditoria(
        obtener_usuario_id_actual(),
        "admin",
        "cambiar_permiso_modulo",
        "usuario",
        user_id,
        {"modulo": modulo, "can_view": can_view, "can_edit": can_edit},
    )
    flash("Permiso de modulo actualizado.", "success")
    return redirect(url_for("admin_usuarios"))


def _resumen_reporte_operativo():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM usuarios")
    usuarios_total = int(cursor.fetchone()[0])
    cursor.execute("SELECT COUNT(*) FROM metas")
    metas_total = int(cursor.fetchone()[0])
    cursor.execute("SELECT COUNT(*) FROM calendario_eventos")
    eventos_total = int(cursor.fetchone()[0])
    cursor.execute("SELECT COUNT(*) FROM mensajes")
    mensajes_total = int(cursor.fetchone()[0])
    cursor.execute("SELECT COUNT(*) FROM audit_log")
    audit_total = int(cursor.fetchone()[0])
    conn.close()
    return {
        "usuarios_total": usuarios_total,
        "metas_total": metas_total,
        "eventos_total": eventos_total,
        "mensajes_total": mensajes_total,
        "audit_total": audit_total,
        "creado_en": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def _reporte_avanzado(f_desde, f_hasta):
    conn = get_connection()
    cursor = conn.cursor()
    where = ["1=1"]
    params = []
    if f_desde:
        where.append("date(creado_en) >= date(?)")
        params.append(f_desde)
    if f_hasta:
        where.append("date(creado_en) <= date(?)")
        params.append(f_hasta)
    sql = f"""
        SELECT modulo, accion, COUNT(*) AS total
        FROM audit_log
        WHERE {' AND '.join(where)}
        GROUP BY modulo, accion
        ORDER BY total DESC, modulo ASC, accion ASC
    """
    cursor.execute(sql, tuple(params))
    rows = cursor.fetchall()
    cursor.execute(
        f"""
        SELECT COUNT(*) AS total
        FROM auth_login_attempts
        WHERE success = 0 AND {' AND '.join(where)}
        """,
        tuple(params),
    )
    login_fail = int((cursor.fetchone() or {"total": 0})["total"])
    conn.close()
    total = sum(int(r["total"] or 0) for r in rows)
    return {"rows": rows, "login_fail": login_fail, "total": total, "desde": f_desde, "hasta": f_hasta}


@app.get("/admin/reportes/excel")
def admin_reporte_excel():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    resumen = _resumen_reporte_operativo()
    salida = StringIO()
    writer = csv.writer(salida)
    writer.writerow(["metrica", "valor"])
    writer.writerow(["creado_en", resumen["creado_en"]])
    writer.writerow(["usuarios_total", resumen["usuarios_total"]])
    writer.writerow(["metas_total", resumen["metas_total"]])
    writer.writerow(["eventos_total", resumen["eventos_total"]])
    writer.writerow(["mensajes_total", resumen["mensajes_total"]])
    writer.writerow(["audit_total", resumen["audit_total"]])

    contenido = salida.getvalue()
    nombre = f"reporte_operativo_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        contenido,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


@app.get("/admin/reportes/pdf")
def admin_reporte_pdf():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    resumen = _resumen_reporte_operativo()
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except Exception:
        flash("Para exportar PDF instala reportlab (pip install reportlab).", "warning")
        return redirect(url_for("admin_dashboard"))

    buff = BytesIO()
    c = canvas.Canvas(buff, pagesize=A4)
    w, h = A4
    y = h - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, "Reporte Operativo - Meta90")
    y -= 24
    c.setFont("Helvetica", 10)
    for k, v in [
        ("Fecha", resumen["creado_en"]),
        ("Usuarios", resumen["usuarios_total"]),
        ("Metas", resumen["metas_total"]),
        ("Eventos", resumen["eventos_total"]),
        ("Mensajes", resumen["mensajes_total"]),
        ("Registros auditoria", resumen["audit_total"]),
    ]:
        c.drawString(40, y, f"{k}: {v}")
        y -= 16
    c.showPage()
    c.save()
    buff.seek(0)
    nombre = f"reporte_operativo_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(
        buff.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


@app.get("/admin/reportes/avanzado")
def admin_reporte_avanzado():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    f_desde = (request.args.get("fecha_desde") or "").strip()
    f_hasta = (request.args.get("fecha_hasta") or "").strip()
    data = _reporte_avanzado(f_desde, f_hasta)
    return render_template("admin/reporte_avanzado.html", data=data)


@app.get("/admin/reportes/avanzado.csv")
def admin_reporte_avanzado_csv():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    f_desde = (request.args.get("fecha_desde") or "").strip()
    f_hasta = (request.args.get("fecha_hasta") or "").strip()
    data = _reporte_avanzado(f_desde, f_hasta)
    out = StringIO()
    w = csv.writer(out)
    w.writerow(["desde", f_desde])
    w.writerow(["hasta", f_hasta])
    w.writerow(["login_fail", data["login_fail"]])
    w.writerow([])
    w.writerow(["modulo", "accion", "total"])
    for r in data["rows"]:
        w.writerow([r["modulo"], r["accion"], r["total"]])
    name = f"reporte_avanzado_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(out.getvalue(), mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename={name}"})


@app.get("/admin/reportes/avanzado.pdf")
def admin_reporte_avanzado_pdf():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    f_desde = (request.args.get("fecha_desde") or "").strip()
    f_hasta = (request.args.get("fecha_hasta") or "").strip()
    data = _reporte_avanzado(f_desde, f_hasta)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except Exception:
        flash("Para exportar PDF instala reportlab (pip install reportlab).", "warning")
        return redirect(url_for("admin_reporte_avanzado"))
    buff = BytesIO()
    c = canvas.Canvas(buff, pagesize=A4)
    w, h = A4
    y = h - 45
    c.setFont("Helvetica-Bold", 13)
    c.drawString(36, y, "Reporte Avanzado Admin")
    y -= 16
    c.setFont("Helvetica", 9)
    c.drawString(36, y, f"Desde: {f_desde or '-'} | Hasta: {f_hasta or '-'} | Login fail: {data['login_fail']}")
    y -= 20
    c.setFont("Helvetica-Bold", 9)
    c.drawString(36, y, "Modulo"); c.drawString(180, y, "Accion"); c.drawString(360, y, "Total")
    y -= 12
    c.setFont("Helvetica", 8)
    for r in data["rows"]:
        if y < 48:
            c.showPage()
            y = h - 45
            c.setFont("Helvetica-Bold", 9)
            c.drawString(36, y, "Modulo"); c.drawString(180, y, "Accion"); c.drawString(360, y, "Total")
            y -= 12
            c.setFont("Helvetica", 8)
        c.drawString(36, y, str(r["modulo"]))
        c.drawString(180, y, str(r["accion"]))
        c.drawString(360, y, str(r["total"]))
        y -= 11
    c.showPage()
    c.save()
    buff.seek(0)
    name = f"reporte_avanzado_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(buff.getvalue(), mimetype="application/pdf", headers={"Content-Disposition": f"attachment; filename={name}"})


@app.get("/admin/operacion/export.zip")
def admin_operacion_export_zip():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    buff = BytesIO()
    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    with zipfile.ZipFile(buff, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        payload = {
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "health": _resumen_health(include_private=True),
            "reporte_diario": _resumen_reporte_diario(),
            "mantenimiento_notificaciones": notificaciones_en_mantenimiento(),
            "tareas": [
                _estado_tarea_windows(r"\Meta90_Notificaciones_Manana"),
                _estado_tarea_windows(r"\Meta90_Notificaciones_Noche"),
                _estado_tarea_windows(r"\Meta90_DB_Backup"),
                _estado_tarea_windows(r"\Meta90_DB_Restore_Verify"),
                _estado_tarea_windows(r"\Meta90_Task_Monitor"),
                _estado_tarea_windows(r"\Meta90_Reporte_Diario"),
            ],
        }
        zf.writestr("estado_operativo.json", json.dumps(payload, ensure_ascii=True, indent=2))

        archivos = [
            os.path.join("logs", "reporte_diario_last.json"),
            os.path.join("logs", "reporte_diario_last.txt"),
            os.path.join("logs", "notificaciones_last.json"),
            os.path.join("logs", "task_monitor.log"),
            os.path.join("logs", "notificaciones_scheduler.log"),
            os.path.join("logs", "db_backup.log"),
            os.path.join("logs", "db_restore_verify.log"),
            os.path.join("logs", "log_rotation.log"),
            os.path.join("logs", "semaforo_alerta.log"),
        ]
        for rel in archivos:
            abs_path = os.path.join(BASE_DIR, rel)
            if os.path.exists(abs_path):
                try:
                    with open(abs_path, "rb") as fh:
                        zf.writestr(rel.replace("\\", "/"), fh.read())
                except OSError:
                    continue

    buff.seek(0)
    return Response(
        buff.getvalue(),
        mimetype="application/zip",
        headers={"Content-Disposition": f"attachment; filename=operacion_meta90_{now}.zip"},
    )


@app.post("/admin/db/restore-emergency")
def admin_db_restore_emergency():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate

    latest = os.path.join(BACKUP_DIR, "metas_latest.db")
    if not os.path.exists(latest):
        flash("No existe backup latest para restaurar.", "warning")
        return redirect(url_for("admin_dashboard"))

    py = os.path.join(BASE_DIR, "venv", "Scripts", "python.exe")
    if not os.path.exists(py):
        py = sys.executable
    restore_cmd = [py, os.path.join(BASE_DIR, "scripts", "db_restore.py"), "--from", latest, "--to", DB_PATH, "--make-safety-copy"]
    verify_cmd = [py, os.path.join(BASE_DIR, "scripts", "db_restore_verify.py")]
    try:
        p1 = subprocess.run(restore_cmd, capture_output=True, text=True, timeout=120, check=False)
        if p1.returncode != 0:
            flash(f"Restore fallo: {(p1.stderr or p1.stdout)[:240]}", "danger")
            return redirect(url_for("admin_dashboard"))
        p2 = subprocess.run(verify_cmd, capture_output=True, text=True, timeout=120, check=False)
        if p2.returncode != 0:
            flash(f"Restore aplicado, pero verify fallo: {(p2.stderr or p2.stdout)[:240]}", "warning")
        else:
            flash("Restore de emergencia completado y verificado.", "success")
        registrar_auditoria(
            obtener_usuario_id_actual(),
            "admin",
            "db_restore_emergency",
            "database",
            None,
            {"backup": latest, "verify_ok": p2.returncode == 0},
        )
    except Exception as exc:  # noqa: BLE001
        flash(f"No se pudo ejecutar restore: {exc}", "danger")
    return redirect(url_for("admin_dashboard"))


@app.get("/health")
def health():
    payload = _resumen_health(include_private=False)
    code = 200 if payload["ok"] else 503
    return jsonify(payload), code


@app.get("/health/details")
def health_details():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    payload = _resumen_health(include_private=True)
    code = 200 if payload["ok"] else 503
    return jsonify(payload), code


@app.get("/api/private/resumen")
def api_private_resumen():
    auth = autenticar_api_token(request.headers.get("Authorization", ""))
    if not auth:
        return jsonify({"ok": False, "error": "unauthorized"}), 401

    user_id = auth["user_id"]
    resumen = obtener_kpis_usuario(user_id)
    return jsonify(
        {
            "ok": True,
            "user": {"id": user_id, "username": auth["username"], "rol": auth["rol"]},
            "resumen": resumen,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    )


@app.get("/mobile")
def mobile_client():
    return render_template("mobile/client.html")


@app.get("/privacy")
def privacy_policy():
    return render_template("legal/privacy.html")


@app.get("/data-deletion")
def data_deletion():
    return render_template("legal/data_deletion.html")


@app.post("/admin/notificaciones/ejecutar")
def admin_ejecutar_notificaciones():
    guard = _admin_guard()
    if guard is not None:
        return guard
    rate = _admin_rate_guard()
    if rate:
        return rate
    if notificaciones_en_mantenimiento():
        return jsonify({"ok": False, "detail": "Notificaciones en mantenimiento."}), 409
    resumen = enviar_recordatorios_sms_todos(forzar_envio=True)
    return jsonify({"ok": True, "resumen": resumen})


configure_perfil_routes(
    {
        "usuario_autenticado": usuario_autenticado,
        "obtener_usuario_id_actual": obtener_usuario_id_actual,
        "get_connection": get_connection,
        "limpiar_telefono": limpiar_telefono,
        "email_valido": email_valido,
        "obtener_prefs_noti": obtener_prefs_notificacion_usuario,
        "guardar_prefs_noti": guardar_prefs_notificacion_usuario,
        "generar_api_token_usuario": generar_api_token_usuario,
        "listar_api_tokens_usuario": listar_api_tokens_usuario,
        "revocar_api_token_usuario": revocar_api_token_usuario,
    }
)
configure_notificaciones_routes(
    {
        "usuario_autenticado": usuario_autenticado,
        "obtener_usuario_id_actual": obtener_usuario_id_actual,
        "obtener_rol_usuario_actual": obtener_rol_usuario_actual,
        "obtener_notificaciones_usuario": obtener_notificaciones_usuario,
        "obtener_estado_sms_operativo": obtener_estado_sms_operativo,
        "sms_configurado": sms_configurado,
        "enviar_recordatorios_sms_usuario": enviar_recordatorios_sms_usuario,
        "push_web_configurado": push_web_configurado,
        "push_public_key": push_public_key,
        "push_web_estado": push_web_estado,
        "guardar_suscripcion_push": guardar_suscripcion_push,
        "enviar_push_usuario": enviar_push_usuario,
        "descartar_notificacion_usuario": descartar_notificacion_usuario,
        "descartar_todas_notificaciones_usuario": descartar_todas_notificaciones_usuario,
        "restaurar_notificaciones_descartadas_hoy": restaurar_notificaciones_descartadas_hoy,
        "aplicar_rate_limit_api": aplicar_rate_limit_api,
    }
)
configure_calendario_routes(
    {
        "calendario": calendario,
        "crear_evento": crear_evento,
        "exportar_eventos_csv": exportar_eventos_csv,
        "exportar_eventos_pdf": exportar_eventos_pdf,
        "importar_eventos": importar_eventos,
        "confirmar_importacion_eventos": confirmar_importacion_eventos,
        "editar_evento": editar_evento,
        "eliminar_evento": eliminar_evento,
    }
)
configure_mensajes_routes(
    {
        "mensajes": mensajes,
        "nuevo_mensaje": nuevo_mensaje,
        "mover_mensaje": mover_mensaje,
        "editar_mensaje": editar_mensaje,
        "eliminar_mensaje": eliminar_mensaje,
    }
)
configure_auth_routes(
    {
        "login": login,
        "admin_2fa": admin_2fa,
        "facebook_start": facebook_start,
        "facebook_callback": facebook_callback,
        "google_start": google_start,
        "google_callback": google_callback,
        "microsoft_start": microsoft_start,
        "microsoft_callback": microsoft_callback,
        "registro": registro,
        "registro_whatsapp": registro_whatsapp,
        "recuperar_cuenta": recuperar_cuenta,
        "restablecer_contrasena": restablecer_contrasena,
        "logout": logout,
    }
)
configure_metas_routes(
    {
        "inicio_metas": inicio_metas,
        "crear_meta": crear_meta,
        "ver_metas": ver_metas,
        "editar": editar,
        "actualizar": actualizar,
        "eliminar": eliminar,
    }
)
app.register_blueprint(auth_bp)
app.register_blueprint(metas_bp)
app.register_blueprint(perfil_bp)
app.register_blueprint(notificaciones_bp)
app.register_blueprint(calendario_bp)
app.register_blueprint(mensajes_bp)


crear_base_datos()


if __name__ == "__main__":
    debug_mode = os.getenv("FLASK_DEBUG", "1").strip() in ("1", "true", "True")
    port = int(os.getenv("PORT", "5000"))
    app.run(debug=debug_mode, host="0.0.0.0", port=port)
